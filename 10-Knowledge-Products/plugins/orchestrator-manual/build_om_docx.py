"""Build OM_Orchestrator_v<VERSION>.docx from OM_Orchestrator_v<VERSION>.md and the PNGs in figures/.

    python3 figures/draw_all.py && python3 build_om_docx.py      # from this folder

VERSION is set below; it is 0.3. Earlier versions are not kept in this folder and are not built
again. The file built is named for the day it was sent to ITU, so that a rendering which was sent
is kept unchanged under a name that says when and to whom.

The .docx is a rendering: it is built, never edited by hand. Every fix goes into the Markdown or
into this script, and the document is built again.

What it makes: an A4 document with a title page, a contents page, the chapters as they stand in
the Markdown (their numbers are in the source), each figure at 6.5 inches wide with its caption,
styled tables, and page numbers. The contents carries page numbers taken from a rendering by
LibreOffice: the script builds once, converts to PDF, finds the page of each heading, and builds
again until the numbers no longer change. Without LibreOffice the contents is built without page
numbers, the script says so, and Word fills them in when the field is updated.

Exit 0 when the document is built; 1 when a check fails (a figure missing, a code line too long).
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
VERSION = "0.3"
SRC = os.path.join(HERE, f"OM_Orchestrator_v{VERSION}.md")
OUT = os.path.join(HERE, f"OM_Orchestrator_v{VERSION}_sent-to-ITU_2026-10-06.docx")

FONT, MONO = "Arial", "Courier New"
NAVY, BLUE, GREY, INK = "1F3864", "2E75B6", "595959", "222222"
TEXT_W = Inches(6.5)            # A4 less 2.25 cm on each side
# Characters a code line may have: 6.5 in is 468 pt, less the box's 0.15 cm indent and its border
# space of 6 pt, leaves about 457 pt; a Courier New character at 8 pt is 4.8 pt wide, so 95 fit.
# One is kept spare. A line of a printed declaration is never broken: the log program reads one
# field per line, so a reader who copies a declaration out of the document must get each field
# whole. (Version 0.2 had 8.5 pt and 88 characters, and broke the declarations to fit.)
CODE_PT = 8
CODE_MAX = 94


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- the Markdown, read into blocks

def read_source(path):
    text = open(path, encoding="utf-8").read()
    meta = {}
    if text.startswith("---\n"):
        end = text.index("\n---\n", 4)
        for line in text[4:end].splitlines():
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
        text = text[end + 5:]
    return meta, text.splitlines()


def blocks(lines):
    """Yield (kind, payload) for each block of the Markdown."""
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "<!-- pagebreak -->":
            yield "pagebreak", None
            i += 1
        elif s.startswith("```"):
            j = i + 1
            code = []
            while j < n and not lines[j].strip().startswith("```"):
                code.append(lines[j].rstrip())
                j += 1
            yield "code", code
            i = j + 1
        elif s.startswith("## "):
            yield "h2", s[3:].strip()
            i += 1
        elif s.startswith("# "):
            yield "h1", s[2:].strip()
            i += 1
        elif s.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            yield "image", (m.group(1), m.group(2))
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            yield "table", rows
        elif s.startswith(">"):
            paras, cur = [], []
            while i < n and lines[i].strip().startswith(">"):
                t = lines[i].strip()[1:].strip()
                if t:
                    cur.append(t)
                elif cur:
                    paras.append(" ".join(cur))
                    cur = []
                i += 1
            if cur:
                paras.append(" ".join(cur))
            yield "quote", paras
        elif re.match(r"^- ", s):
            items = []
            while i < n and re.match(r"^- ", lines[i].strip()):
                items.append(lines[i].strip()[2:])
                i += 1
            yield "bullets", items
        elif re.match(r"^\d+\. ", s):
            items = []
            while i < n and re.match(r"^\d+\. ", lines[i].strip()):
                m = re.match(r"^(\d+)\. (.*)$", lines[i].strip())
                items.append((m.group(1), m.group(2)))
                i += 1
            yield "numbers", items
        else:
            para = []
            while i < n and lines[i].strip() and not re.match(
                    r"^(#|```|!\[|\||>|- |\d+\. |<!--)", lines[i].strip()):
                para.append(lines[i].strip())
                i += 1
            p = " ".join(para)
            if re.match(r"^\*Figure \d+\..*\*$", p):
                yield "caption", p[1:-1]
            else:
                yield "para", p


# ---------------------------------------------------------------- inline text

def smart(t):
    """Straight quotes to typographic ones, outside code."""
    t = re.sub(r'(^|[\s(\[\u2014])"', "\\1\u201c", t)
    t = t.replace('"', "\u201d")
    t = re.sub(r"(\w)'(\w)", "\\1\u2019\\2", t)
    t = re.sub(r"(^|[\s(\[])'", "\\1\u2018", t)
    t = t.replace("'", "\u2019")
    return t


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


def add_inline(p, text, size=None, bold=False, italic=False, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(p, part[2:-2], size, True, italic, color)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt((size or 10.5) - 0.5)
            r.font.color.rgb = rgb("1B2A31")
            r.bold, r.italic = (True if bold else None), (True if italic else None)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(p, part[1:-1], size, bold, not italic, color)
        else:
            r = p.add_run(smart(part))
            r.bold, r.italic = (True if bold else None), (True if italic else None)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)


def plain(text):
    """The text of a Markdown span without its marks, as it will read in the document."""
    return smart(re.sub(r"\*\*|`|\*", "", text))


# ---------------------------------------------------------------- the document's styles

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para_border(p, side, color, sz=6, space=4):
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def para_shading(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def style_font(style, size, bold=False, italic=False, color=INK, name=FONT):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-GB")


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.25)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.0)
    sec.different_first_page_header_footer = True

    st = doc.styles
    # the document's default font is Arial too, not the template's theme font
    dd = st.element.find(qn("w:docDefaults"))
    rf = dd.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts")) if dd is not None else None
    if rf is not None:
        for a in list(rf.attrib):
            del rf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
    normal = st["Normal"]
    style_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    h1 = st["Heading 1"]
    style_font(h1, 18, bold=True, color=NAVY)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.keep_with_next = True
    para_border_style(h1, "bottom", BLUE)

    h2 = st["Heading 2"]
    style_font(h2, 13, bold=True, color=BLUE)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    def new(name, base="Normal"):
        s = st.add_style(name, 1)
        s.base_style = st[base]
        return s

    cap = new("OM Caption")
    style_font(cap, 9, italic=True, color=GREY)
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(12)

    fig = new("OM Figure")
    fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_before = Pt(6)
    fig.paragraph_format.space_after = Pt(0)
    fig.paragraph_format.keep_with_next = True
    fig.paragraph_format.line_spacing = 1.0

    code = new("OM Code")
    style_font(code, CODE_PT, name=MONO, color="1B2A31")
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.15)
    code.paragraph_format.keep_together = True

    q = new("OM Quote")
    style_font(q, 10, color="263238")
    q.paragraph_format.left_indent = Cm(0.6)
    q.paragraph_format.right_indent = Cm(0.3)
    q.paragraph_format.space_after = Pt(4)

    for name, ind, hang in (("OM Bullet", 0.6, 0.45), ("OM Number", 0.75, 0.65)):
        s = new(name)
        s.paragraph_format.left_indent = Cm(ind)
        s.paragraph_format.first_line_indent = Cm(-hang)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.tab_stops.add_tab_stop(Cm(ind))

    t1 = new("OM TOC 1")
    style_font(t1, 10.5, bold=True, color=NAVY)
    t1.paragraph_format.space_before = Pt(7)
    t1.paragraph_format.space_after = Pt(1)
    t1.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    t2 = new("OM TOC 2")
    style_font(t2, 10)
    t2.paragraph_format.left_indent = Cm(0.7)
    t2.paragraph_format.space_after = Pt(0)
    t2.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    cell = new("OM Cell")
    style_font(cell, 9.5)
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.line_spacing = 1.05

    # header and footer: nothing on the title page, a running title and the page number after it
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("The Orchestrator \u00b7 an operational manual")
    r.font.size, r.font.color.rgb, r.font.name = Pt(8.5), rgb(GREY), FONT
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.size, r.font.color.rgb = Pt(9), rgb(GREY)
    add_field(fp, "PAGE", "1", size=9)


def para_border_style(style, side, color):
    pPr = style.element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "8")
    el.set(qn("w:space"), "6")
    el.set(qn("w:color"), color)
    pbdr.append(el)
    pPr.append(pbdr)


def fld(run, kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def add_field(p, instr, result, size=None):
    fld(p.add_run(), "begin")
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld(p.add_run(), "separate")
    r = p.add_run(result)
    if size:
        r.font.size, r.font.color.rgb = Pt(size), rgb(GREY)
    fld(p.add_run(), "end")


# ---------------------------------------------------------------- the parts of the document

def title_page(doc, meta):
    def line(text, size, bold=False, italic=False, color=INK, before=0, after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.size, r.bold, r.italic, r.font.color.rgb = Pt(size), bold, italic, rgb(color)
        return p

    line(meta.get("series", ""), 11, color=GREY, before=150, after=24)
    line(meta.get("title", ""), 40, bold=True, color=NAVY, after=10)
    p = line(meta.get("subtitle", ""), 16, color=BLUE, after=18)
    para_border(p, "bottom", BLUE, sz=8, space=12)
    line(meta.get("question", ""), 15, italic=True, color=INK, before=6, after=200)
    line(meta.get("edition", ""), 10, color=GREY)


def contents(doc, entries):
    """A real TOC field, with the entries filled in as its result, so that the contents reads
    correctly before anyone updates it and Word can still update it."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Contents")
    r.font.size, r.bold, r.font.color.rgb = Pt(18), True, rgb(NAVY)
    para_border(p, "bottom", BLUE, sz=8, space=6)
    last = len(entries) - 1
    for k, (level, text, page) in enumerate(entries):
        q = doc.add_paragraph(style="OM TOC 1" if level == 1 else "OM TOC 2")
        if k == 0:
            fld(q.add_run(), "begin")
            r = q.add_run()
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-2" \\h \\z \\u '
            r._r.append(it)
            fld(q.add_run(), "separate")
        q.add_run(text)
        q.add_run("\t" + (str(page) if page else ""))
        if k == last:
            fld(q.add_run(), "end")


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    ncol = len(header)
    minw, avg = [], []
    for j in range(ncol):
        col = [plain(r[j]) if j < len(r) else "" for r in [header] + body]
        word = max(len(w) for c in col for w in (c.split() or [""]))
        minw.append(word * 0.078 + 0.22)
        avg.append(sum(len(c) for c in col[1:]) / max(1, len(col) - 1) + 4)
    total_in = TEXT_W.inches
    if sum(minw) > total_in:
        minw = [w * total_in / sum(minw) for w in minw]
    spare = total_in - sum(minw)
    widths_in = [minw[j] + spare * avg[j] / sum(avg) for j in range(ncol)]
    widths = [Inches(w) for w in widths_in]
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 45), ("bottom", 45), ("left", 90), ("right", 90)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    for i, row in enumerate([header] + body):
        tr = t.rows[i]
        trPr = tr._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        if i == 0:
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        for j in range(ncol):
            c = tr.cells[j]
            c.width = widths[j]
            p = c.paragraphs[0]
            p.style = doc.styles["OM Cell"]
            add_inline(p, row[j] if j < len(row) else "", size=9.5, bold=(i == 0),
                       color=NAVY if i == 0 else None)
            if i == 0:
                set_cell_shading(c, "F2F2F2")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(meta, lines, toc_pages, out):
    doc = Document()
    setup(doc)
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.subject = meta.get("subtitle", "")
    doc.core_properties.author = "ITU/Giga Knowledge Products"
    doc.core_properties.last_modified_by = "ITU/Giga Knowledge Products"
    doc.core_properties.comments = ""
    doc.core_properties.keywords = ""

    items = list(blocks(lines))
    heads = [(1 if k == "h1" else 2, plain(v)) for k, v in items if k in ("h1", "h2")]
    title_page(doc, meta)
    contents(doc, [(lv, tx, toc_pages.get((lv, tx))) for lv, tx in heads])

    figures, tables, problems = 0, 0, []
    breaknext = False
    for idx, (kind, v) in enumerate(items):
        nxt = items[idx + 1][0] if idx + 1 < len(items) else None
        if kind == "pagebreak":
            breaknext = True
            continue
        if kind == "h1":
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, v)
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, v)
        elif kind == "para":
            p = doc.add_paragraph()
            add_inline(p, v)
            label = re.match(r"^\*\*[^*]+\*\*$", v) is not None
            if label or nxt in ("bullets", "numbers", "code", "table", "quote", "image"):
                p.paragraph_format.keep_with_next = True
        elif kind == "bullets":
            for it in v:
                p = doc.add_paragraph(style="OM Bullet")
                p.add_run("\u2022\t")
                add_inline(p, it)
            p.paragraph_format.space_after = Pt(6)
        elif kind == "numbers":
            for num, it in v:
                p = doc.add_paragraph(style="OM Number")
                p.add_run(f"{num}.\t")
                add_inline(p, it)
            p.paragraph_format.space_after = Pt(6)
        elif kind == "quote":
            for k, it in enumerate(v):
                p = doc.add_paragraph(style="OM Quote")
                para_border(p, "left", BLUE, sz=18, space=8)
                para_shading(p, "F4F7FB")
                add_inline(p, it, size=10)
                if k < len(v) - 1:
                    p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(10)
        elif kind == "code":
            p = doc.add_paragraph(style="OM Code")
            para_shading(p, "F5F5F5")
            para_border(p, "left", "90A4AE", sz=12, space=6)
            for k, ln in enumerate(v):
                if len(ln) > CODE_MAX:
                    problems.append(f"code line of {len(ln)} characters: {ln[:50]}")
                r = p.add_run(ln if ln else " ")
                if k < len(v) - 1:
                    r.add_break()
        elif kind == "image":
            alt, src = v
            path = os.path.join(HERE, src)
            if not os.path.exists(path):
                problems.append(f"missing figure {src}")
                continue
            p = doc.add_paragraph(style="OM Figure")
            r = p.add_run()
            r.add_picture(path, width=TEXT_W)
            for d in r._r.iter(qn("wp:docPr")):
                d.set("descr", alt)
                d.set("title", alt)
            figures += 1
        elif kind == "caption":
            p = doc.add_paragraph(style="OM Caption")
            m = re.match(r"^(Figure \d+\.)(.*)$", v)
            r = p.add_run(m.group(1))
            r.bold = True
            add_inline(p, m.group(2), size=9, italic=True)
        elif kind == "table":
            add_table(doc, v)
            tables += 1
            p = None
        else:
            continue
        if breaknext and kind not in ("table",) and p is not None:
            p.paragraph_format.page_break_before = True
            breaknext = False
    doc.save(out)
    return heads, figures, tables, problems


# ---------------------------------------------------------------- page numbers for the contents

def page_of_heads(docx, heads, work):
    """Convert with LibreOffice, read the text of each page, and find the page of each heading."""
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", work, docx],
                       capture_output=True, text=True, timeout=600)
    pdf = os.path.join(work, os.path.splitext(os.path.basename(docx))[0] + ".pdf")
    if r.returncode != 0 or not os.path.exists(pdf):
        raise RuntimeError(f"soffice failed: {r.stderr.strip()[:300]}")
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True,
                         check=True).stdout
    pages = txt.split("\f")

    def norm(s):
        s = s.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'")
        return re.sub(r"\s+", " ", s).strip()

    clean = []
    for pg in pages:
        keep = [ln for ln in pg.splitlines() if not re.search(r"\.{4,}\s*\d*\s*$", ln)]
        clean.append(norm(" ".join(keep)))
    found, cursor = {}, 0
    for lv, tx in heads:
        target = norm(tx)
        for k in range(cursor, len(clean)):
            if target in clean[k]:
                found[(lv, tx)] = k + 1
                cursor = k
                break
    return found, len([p for p in pages if p.strip()]), pdf


def main():
    meta, lines = read_source(SRC)
    have_lo = shutil.which("soffice") and shutil.which("pdftotext")
    toc = {}
    work = tempfile.mkdtemp(prefix="om_build_")
    try:
        heads, figures, tables, problems = build(meta, lines, toc, OUT)
        if problems:
            print("\n".join(problems))
            return 1
        if not have_lo:
            print("LibreOffice or pdftotext not found: the contents has no page numbers; "
                  "update the field in Word")
        else:
            for attempt in range(1, 5):
                pages, npages, _ = page_of_heads(OUT, heads, work)
                missing = [tx for lv, tx in heads if (lv, tx) not in pages]
                if missing:
                    print(f"headings not found in the rendering: {missing}")
                    return 1
                if pages == toc:
                    break
                toc = pages
                build(meta, lines, toc, OUT)
            qr = toc.get((1, "Quick reference"))
            gl = toc.get((1, "Glossary"))
            print(f"pages: {npages}; contents settled after {attempt} rendering(s); "
                  f"quick reference: pages {qr} to {gl - 1} ({gl - qr} pages)")
        print(f"built {OUT}: {figures} figures, {tables} tables, {len(heads)} headings")
        return 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
