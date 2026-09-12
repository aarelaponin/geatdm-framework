#!/usr/bin/env python3
"""Did the take actually cover every slide, or did it skip some?

Runtime is a bad proxy. A take can be short because the hosts were efficient, or short because
they silently dropped three slides — and the drift check cannot tell those apart. This can: for
each slide it takes the terms that are DISTINCTIVE to that slide (present in its voice-over,
rare elsewhere in the deck) and looks for them in the transcript.

    python3 coverage_check.py <deck.pptx> <take.srt>

THIN is a prompt to read the transcript, not a verdict. The method matches vocabulary, so a
segment the hosts covered thoroughly in their own words scores low: 1.7's Ask 4 read 28% while
the take actually said "a written promise that this EA team will not be pulled onto the urgent
project of the week". MISS on a slide with many distinctive terms is reliable; THIN on a short
slide (few terms, so one word swings the percentage) usually is not. Check before re-rolling.
"""
import re
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

STOP = set("""the a an and or but if then that this these those of to in on for with as by at
from is are was were be been being it its他 they them their you your our we us have has had do
does did not no so than there here what which who whom whose when where why how all any each
every some most more much many one two three four five will would can could should may might
must into out up down over under again further once about against between through during before
after above below only own same such nor too very just now also make makes made take takes
because while what's it's don't doesn't isn't""".split())
WORD = re.compile(r"[a-z][a-z-]{3,}")

# Deck furniture is on the slide but is never spoken, so counting it as content marks a
# perfectly covered section slide "thin" for not saying "voice-over" and "standalone".
CHROME = re.compile(
    r"^(?:\d\.\d\s*·|KP\d+\s*·\s*Government|~\d+\s*minutes?\s*·|Length:|Target audience:|"
    r"www\.itu\.int|IN ONE SENTENCE|WHERE WE START|THE LIFECYCLE THIS MODULE TEACHES|\d+\s*sign-offs?\s*·)",
    re.I)


def slide_terms(prs):
    """-> [(slide_no, title, {distinctive terms})]"""
    per, titles = [], []
    for i, sl in enumerate(prs.slides, 1):
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        vo = " ".join(m.group(1) for line in notes.split("\n")
                      for m in [re.match(r"^VO(?:, slide \d+)?:\s*(.*)$", line.strip())] if m)
        vis = " ".join(
            "\n".join(l for l in sh.text_frame.text.splitlines()
                      if l.strip() and not CHROME.match(l.strip()))
            for sh in sl.shapes
            if sh.has_text_frame and not sh.text_frame.text.startswith("Do this on"))
        t = next((x.strip() for x in vis.splitlines() if x.strip()), f"slide {i}")
        per.append(set(WORD.findall((vo + " " + vis).lower())) - STOP)
        titles.append(t[:52])
    df = Counter(w for s in per for w in s)
    # distinctive = appears on this slide and at most one other
    return [(i + 1, titles[i], {w for w in s if df[w] <= 2}) for i, s in enumerate(per)]


def advisory_slides(prs):
    """The two slides whose vocabulary the narration is not expected to reuse.

    WHERE WE START is a ~40-word scripted opener the hosts stretch to 45-90 s in their own
    words, usually behind an analogy of their own. IN ONE SENTENCE restates the content slides,
    so whatever vocabulary is *distinctive* to it is whatever the others did not use — on 5.6
    that is "held", "possible" and "whole", three function words. Neither measures coverage.
    """
    out = set()
    for i, sl in enumerate(prs.slides, 1):
        text = " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
        if re.search(r"WHERE WE START|IN ONE SENTENCE", text, re.I):
            out.add(i)
    return out


def main():
    prs = Presentation(sys.argv[1])
    srt = Path(sys.argv[2]).read_text(encoding="utf-8-sig")
    take = set(WORD.findall(" ".join(
        " ".join(b.splitlines()[2:]) for b in re.split(r"\n\s*\n", srt.strip())
        if len(b.splitlines()) >= 3).lower()))

    print(f"=== coverage — {Path(sys.argv[2]).name} against {Path(sys.argv[1]).name}\n")
    thin = 0
    slides = slide_terms(prs)
    advisory = advisory_slides(prs)
    # the title card and the Sources slide are bookends: the first is a 15 s cold open and
    # nothing is spoken over the last, so neither is a coverage failure
    for n, title, terms in slides[1:-1]:
        if not terms:
            print(f"  --   slide {n:>2}  {title}  (no distinctive terms — skipped)")
            continue
        hit = terms & take
        pct = len(hit) / len(terms)
        flag = "ok  " if pct >= 0.34 else ("THIN" if pct >= 0.15 else "MISS")
        # The opener and the recap join the title card and the Sources slide as bookends —
        # see advisory_slides(). Nine of the twenty-three shipped KP1 takes score 17-33% on the
        # opener. Reported, never a failure.
        if n in advisory and flag != "ok  ":
            flag = "note"
        thin += flag not in ("ok  ", "note")
        print(f"  {flag} slide {n:>2}  {title}  {len(hit)}/{len(terms)} terms ({pct:.0%})")
        if flag not in ("ok  ", "note"):
            print(f"         missing: {', '.join(sorted(terms - take)[:10])}")
    print(f"\n{thin} slide(s) thin or missing")
    return 1 if thin else 0


if __name__ == "__main__":
    sys.exit(main())
