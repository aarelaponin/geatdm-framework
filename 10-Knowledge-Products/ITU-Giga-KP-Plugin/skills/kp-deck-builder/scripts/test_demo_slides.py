#!/usr/bin/env python3
"""The three demo-evidence slide kinds, built from fixture takes.

What has to hold: the notes still carry the VO untouched (vo_diff.py, scripts_from_deck.py) with a
clip slide's CLIP: line in front; every slide carries the provenance line from takes.json; a frame
is letterboxed — aspect kept, centred, inside the content area. Run: python3 test_demo_slides.py
"""
import json
import os
import tempfile

from PIL import Image
from pptx.util import Emu

import deck_lib as dl

VO = 'VO: The form has one field typed and nine filled in.'


def build(tmp):
    json.dump({'xroad_version': '7.7.0', 'pack_commit': '1a2b3c4d5e6f', 'captured_at': '2026-09-13T15:00:00Z',
               'beats': {}}, open(os.path.join(tmp, 'takes.json'), 'w'))
    Image.new('RGB', (1920, 1080), 'white').save(os.path.join(tmp, 'C2.png'))
    open(os.path.join(tmp, 'C8.txt'), 'w').write('\n'.join('2.6.%d  PASS  check %d' % (i, i) for i in range(20)))
    open(os.path.join(tmp, 'REQ.md'), 'w').write('# Requirements\n\n| Requirement | Stated |\n|---|---|\n'
                                                  '| Security server | yes |\n| Lawful basis | stated |\n')
    prs = dl.open_template()
    dl.delete_template_slides(prs, keep=0)
    dl.demo_slide(prs, 'The form after', os.path.join(tmp, 'C2.png'), 'nine fields not asked', 'tag', VO,
                  clip=os.path.join(tmp, 'C6.mp4'))
    dl.terminal_slide(prs, 'The acceptance check', os.path.join(tmp, 'C8.txt'), 'six checks green', 'tag', VO)
    dl.artefact_slide(prs, 'The requirements record', os.path.join(tmp, 'REQ.md'), 'as stated', 'tag', VO)
    return prs


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_notes_keep_the_vo_and_a_clip_leads_with_its_file():
    with tempfile.TemporaryDirectory() as tmp:
        frame, term, art = build(tmp).slides
        assert frame.notes_slide.notes_text_frame.text == 'CLIP: C6.mp4\n\n' + VO
        assert term.notes_slide.notes_text_frame.text == VO
        assert art.notes_slide.notes_text_frame.text == VO


def test_every_evidence_slide_has_its_provenance_and_marker():
    want = 'captured from the demonstration federation · X-Road 7.7.0 · pack 1a2b3c4 · 2026-09-13'
    with tempfile.TemporaryDirectory() as tmp:
        for s in build(tmp).slides:
            assert want in texts(s), texts(s)
            assert any(sh.name == dl.DEMO_SHAPE for sh in s.shapes)


def test_frame_is_letterboxed_inside_the_content_area():
    with tempfile.TemporaryDirectory() as tmp:
        pic = next(sh for sh in build(tmp).slides[0].shapes if sh.shape_type == 13)
        x, y, w, h = (Emu(v).inches for v in (pic.left, pic.top, pic.width, pic.height))
        ax, ay, aw, ah = dl.DEMO_AREA
        assert abs(w / h - 16 / 9) < 0.01, w / h
        assert abs(h - ah) < 0.01                        # height-bound for a 16:9 frame
        assert abs((x - ax) - (ax + aw - x - w)) < 0.01  # centred
        assert ax <= x and x + w <= ax + aw + 0.001


def test_a_long_capture_is_cut_not_wrapped():
    with tempfile.TemporaryDirectory() as tmp:
        term = build(tmp).slides[1]
        block = next(t for t in texts(term) if t.startswith('2.6.0'))
        lines = block.split('\n')
        assert len(lines) == 18 and lines[-1] == '…', lines[-3:]


if __name__ == '__main__':
    for name, fn in sorted(globals().items()):
        if name.startswith('test_'):
            fn()
            print(f'ok  {name}')
    print('all ok')
