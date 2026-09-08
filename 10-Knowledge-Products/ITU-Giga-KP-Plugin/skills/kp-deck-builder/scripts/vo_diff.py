#!/usr/bin/env python3
"""vo_diff.py — prove a module deck narrates exactly its build script (plan §2, WP7).

Usage: vo_diff.py build_kp1_moduleN_v0X.js KP1_MN_Deck_v0.Y.pptx [--stats]

The .js is the source of truth for the deliverable; the deck's notes are the source of truth
for what gets said. This reports every `scriptBeats[].text` that does not appear verbatim in
the notes of that subtopic's slides, and every `VO:` note paragraph that is not in the beats.
Also prints the per-video word/slide table and the practice-box / handoff checks.
Exit 1 on any mismatch. Whitespace, curly quotes and dashes are normalised."""
import re, sys
from pptx import Presentation

def norm(s):
    for a, b in (('’', "'"), ('‘', "'"), ('“', '"'), ('”', '"'), ('—', '-'), ('–', '-'),
                 ("\\'", "'"), ('\\"', '"'), (' ', ' ')):
        s = s.replace(a, b)
    s = re.sub(r'''["']''', '', s)  # quote style differs between .js escapes and notes
    return re.sub(r'\s+', ' ', s).strip().lower()

def js_beats(path):
    js = open(path, encoding='utf-8').read()
    marks = [(m.group(1), m.start()) for m in re.finditer(r'num:\s*"\d+\.\d+ Subtopic (\d+\.\d+)"', js)]
    marks.append(('END', len(js)))
    out = {}
    for i in range(len(marks) - 1):
        code, a = marks[i]; b = marks[i + 1][1]
        blk = js[a:b]
        sb = blk[blk.find('scriptBeats'):]
        if 'slideSpecRows' in sb: sb = sb[:sb.find('slideSpecRows')]
        texts = re.findall(r'\{\s*text:\s*"((?:[^"\\]|\\.)*)"', sb)
        cues = re.findall(r'\{\s*cue:\s*"((?:[^"\\]|\\.)*)"', sb)
        out[code] = ([norm(t) for t in texts], len(cues))
    return out

def deck_sections(path):
    prs = Presentation(path)
    secs, cur = {}, None
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame:
                m = re.match(r'KP\d · MODULE \d · VIDEO (\d+\.\d+)$', sh.text_frame.text.strip())
                if m: cur = m.group(1); secs[cur] = []
        if cur is None: continue
        if s.slide_layout.name.strip().startswith('2_Thank'): break
        vo, has_box, says_box = [], False, False
        if s.has_notes_slide:
            for p in s.notes_slide.notes_text_frame.text.split('\n'):
                m = re.match(r'VO(?:,\s*slide\s*\d+)?:\s*(.*)', p.strip())
                if m:
                    vo.append(norm(m.group(1)))
                    if re.search(r'do this on your own sector|in the description|companion material', vo[-1]): says_box = True
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith('Do this on your own sector'): has_box = True
        secs[cur].append((i, vo, has_box, says_box))
    return secs

def main():
    if len(sys.argv) < 3: sys.exit(__doc__)
    beats = js_beats(sys.argv[1]); secs = deck_sections(sys.argv[2])
    bad = 0; total = 0
    print('| video | slides | VO words | js words | beats | mismatches | box |')
    print('|---|---|---|---|---|---|---|')
    details = []
    for code, (texts, ncues) in beats.items():
        if code not in secs:
            print(f'| {code} | — | — | — | {len(texts)} | NOT IN DECK | — |'); bad += 1; continue
        sl = secs[code]
        dparas = [p for _, vo, _, _ in sl for p in vo]
        dtext = ' '.join(dparas); jtext = ' '.join(texts)
        miss = [t for t in texts if t not in dtext]
        extra = [d for d in dparas if d not in jtext]
        boxes = sum(1 for _, _, b, _ in sl if b); says = [i for i, _, _, sb in sl if sb]
        nm = len(miss) + len(extra); bad += nm
        if boxes != 1: bad += 1
        if says: bad += 1
        dw = len(dtext.split()); total += dw
        print(f'| {code} | {len(sl)} ({sl[0][0]}–{sl[-1][0]}) | {dw} | {len(jtext.split())} | {len(texts)}/{ncues} cues | {nm} | {boxes}{" NARRATED@"+str(says) if says else ""} |')
        for t in miss: details.append(f'  {code} JS beat not in deck notes: {t[:140]}…')
        for d in extra: details.append(f'  {code} deck VO not in .js:        {d[:140]}…')
    for code in secs:
        if code not in beats: print(f'| {code} | in deck, NOT IN JS |'); bad += 1
    print(f'\nTotal deck VO words: {total}')
    if details: print('\n'.join(details))
    print('\nRESULT:', 'OK — zero mismatches, one practice box per video, none narrated' if not bad else f'{bad} problem(s)')
    sys.exit(1 if bad else 0)

if __name__ == '__main__': main()
