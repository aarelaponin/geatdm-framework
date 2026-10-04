#!/usr/bin/env python3
# Build the KP3 Module 1 video deck on the ITU template — v0.1.
# Content follows KP3_Module1_Script_Bundle_v0.2 (build_kp3_module1_v02.js): ten videos,
# 1.1 – 1.10, Strategist-facing — where your country stands, and what to build first. Every VO
# paragraph in the notes is a verbatim scriptBeats[].text from the .js (vo_diff.py proves it);
# the recap slide of every video carries the single message word for word and the un-narrated
# practice box (plan D5) — task = the AI tip's title, artefact = the subtopic's `practice` field.
# The demonstration segments of 1.4 and 1.7 (not yet recorded) are shown as their storyboard's
# steps on a text slide marked 'WHAT A GOOD RUN SHOWS', as the bundle's §4.6 asks.
# Content only — every generic helper, branding constant and layout index comes from
# $KP_KIT/skills/kp-deck-builder/scripts/deck_lib.py (which also ships the
# template). Conventions and design rules: that skill's SKILL.md. The three content-shaped
# helpers below (table_slide, storyboard_slide, recap) compose deck_lib primitives only.
# Generated .pptx is NEVER hand-edited — fix here, re-render, re-run the split
# (kp-deck-builder/scripts/split_module_deck.py + the split spec next to the decks).
# Override paths with TEMPLATE= and OUT_PATH= env vars.
# 4 Oct 2026: 1.8's registers slide reworded (title, rows, VO) so the narration does not hang on
# the name PAERA — nine NotebookLM takes skipped the slide; citations stay on the Sources card.
import os
import sys

KP_KIT = os.environ.get('KP_KIT') or os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'ITU-Giga-KP-Plugin')
if not os.path.isdir(os.path.join(KP_KIT, 'skills', 'kp-deck-builder', 'scripts')):
    sys.exit("Set KP_KIT to the itu-giga-kp folder of your claude-marketplace clone (plugins/itu-giga-kp).")
sys.path.insert(0, os.path.join(KP_KIT, 'skills', 'kp-deck-builder', 'scripts'))
from pptx.util import Pt
from deck_lib import (
    TITLE_CARD_NOTE, hook_slide, practice_box,
    GREY, INK, ITU_BLUE, ITU_BLUE_DARK, LIGHT, PANEL_GREY, WHITE,
    LAYOUT_BLUE, LAYOUT_THANKS, LAYOUT_WHITE,
    add_slide, big_slide, block_slide, box, delete_template_slides, edit_agenda,
    edit_cover, footer, hline, notes, open_template, rows_block, rows_slide, section_slide,
    set_text, sources_slide, title, two_panel)
from deck_diagrams import node

prs = open_template(os.environ.get('TEMPLATE'))

AUDIENCE = ('The public-sector middle manager who commissions the assessment, defends its '
            'result to the minister and plans what to build first')

# The practice box is the video's only call to action and is never narrated (plan D5).
PRACTICE_NOTE = ('PRACTICE BOX (on-screen only — never read it, never paraphrase it, never point '
                 'at it). It replaces the narrated handoff: this video ends on the recap and the '
                 'Sources slide.')
STORYBOARD_LABEL = 'WHAT A GOOD RUN SHOWS'
STORYBOARD_NOTE = ('On screen: the storyboard\'s steps as text, marked \'What a good run shows\', in the place of the demonstration segment until it is recorded. The '
                   'narration says what the walkthrough shows and what counts as a pass, never '
                   'that it has run.')


def block(prs, *a, **k):
    k.setdefault('punch_y', 4.35)
    return block_slide(prs, *a, **k)


def panels(prs, *a, **k):
    k.setdefault('height', 3.2)
    return two_panel(prs, *a, **k)


# Opener (hook) slide copy per video — headline + two to four supporting lines, written from
# the same opener narration the hook's note carries. Not a preview of the next slide's list.
HOOKS = {'1.1': ('Donors and vendors all speak of DPI.',
                 ['Some of what they offer is shared by the whole government. Some of it is one '
                  'ministry\'s system under a new name.',
                  'Before your ministry pays for either, you need a test you can defend.']),
         '1.2': ('Ask a ministry what the country has, and you get a list of projects.',
                 ['Each with its own donor. A list hides what is missing.',
                  'One page with five domains shows your minister both.']),
         '1.3': ('A minister who commissions a roadmap will ask three things.',
                 ['Who does the work? When must I decide? How will anyone know the result is '
                  'right?',
                  'Nine named steps let you answer all three before the work starts.']),
         '1.4': ('An assessment usually starts by asking officials for documents — and then '
                 'waiting.',
                 ['Start instead with what your country has already published.',
                  'In a few days an AI assistant gives you a first picture, and a list of points '
                  'to check.']),
         '1.5': ('The public record tells you what a country says about its systems.',
                 ['The people who run them know what the systems actually do.',
                  'Ask them, in writing — and ask for the document behind every answer.']),
         '1.6': ('A questionnaire answer is what a body says about itself.',
                 ['A desk result is what the public record suggests. Neither is a finding yet.',
                  'Score them as they stand, and the first official who disagrees can overturn '
                  'your result.']),
         '1.7': ('Two assessors who read the same evidence should reach the same score.',
                 ['If they do not, the score is an opinion, and a minister can set it aside.',
                  'Written criteria on a published scale prevent that.']),
         '1.8': ('When a ministry asks for money for a system, the first question is whose '
                 'system it is.',
                 ['Shared by every sector — or one sector\'s own?',
                  'Get that wrong, and the budget pays for the same thing twice.']),
         '1.9': ('Money for digital government comes in pieces.',
                 ['A donor here, a budget line there. Each piece is spent on something.',
                  'Unless someone has fixed the order of building, the pieces arrive in an order '
                  'that nobody chose.']),
         '1.10': ('A roadmap that promises everything proves nothing for years.',
                  ['Pick one service that shows the shared blocks serving real people, and prove '
                   'it first.',
                   'For education, that service is registering a learner once.'])}


def section(code, name, message, note):
    # No runtime on the title card: the narration is generated per take and its length moves
    # with every re-roll.
    s = section_slide(prs, 'MODULE 1 · VIDEO %s' % code, code, name, message,
                      'standalone video · voice-over on text slides', TITLE_CARD_NOTE)
    head, lines = HOOKS[code]
    hook_slide(prs, head, lines, '%s · %s' % (code, name), note)
    return s


def closing_line(s, text, y, size=15.5):
    tb = box(s, 0.72, y, 11.9, 0.6)
    set_text(tb.text_frame, [[(text, size, True, ITU_BLUE_DARK, False)]])


def foot_line(s, text, y=6.35):
    """The small grey line under a worked example: 'Illustrative. Worked example E5.'"""
    tb = box(s, 0.72, y, 11.9, 0.4)
    set_text(tb.text_frame, [[(text, 12, False, GREY, True)]])


def table_slide(prs, head, cols, widths, rows, tag, note, foot=None, size=13, top=1.55,
                bottom=6.25):
    """A plain text table — the shape the bundle's 'N-row text table' cues ask for. Header row
    in the accent colour, thin separators, no fills: text only, per the ITU guide."""
    s = add_slide(prs, LAYOUT_WHITE)
    title(s, head)
    xs = [0.72]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    hh = 0.42
    for x, w, c in zip(xs, widths, cols):
        tb = box(s, x, top, w - 0.12, hh)
        set_text(tb.text_frame, [[(c, size, True, ITU_BLUE_DARK, False)]])
    hline(s, 0.72, top + hh, 11.9, color=ITU_BLUE_DARK, weight_pt=1.25)
    rh = (bottom - top - hh - 0.08) / len(rows)
    for i, row in enumerate(rows):
        y = top + hh + 0.08 + i * rh
        for j, (x, w, cell) in enumerate(zip(xs, widths, row)):
            tb = box(s, x, y + 0.04, w - 0.12, rh - 0.08)
            set_text(tb.text_frame, [[(cell, size, j == 0, INK, False)]])
        if i < len(rows) - 1:
            hline(s, 0.72, y + rh - 0.02, 11.9)
    if foot:
        foot_line(s, foot, bottom + 0.1)
    footer(s, tag)
    notes(s, note)
    return s


def storyboard_slide(prs, head, steps, tag, note):
    """The text stand-in for a demonstration segment that has not been recorded (bundle §4.6):
    the storyboard's steps as numbered rows under a plain label."""
    s = add_slide(prs, LAYOUT_WHITE)
    title(s, head)
    tb = box(s, 0.72, 1.4, 6, 0.35)
    set_text(tb.text_frame, [[(STORYBOARD_LABEL, 12, True, GREY, False)]])
    rows_slide(s, [(st, '') for st in steps], top=1.85, bottom=6.5, numbered=True, head_size=17)
    footer(s, tag)
    notes(s, note)
    return s


def recap(text, tag, note, practice):
    """big_slide, with the type stepped down for the longer single messages so the line stays
    clear of the practice box (the bundle carries each message word for word)."""
    if len(text) <= 150:
        return big_slide(prs, text, tag, note, practice=practice)
    s = add_slide(prs, LAYOUT_BLUE)
    tb = box(s, 1.1, 1.7, 6, 0.4)
    set_text(tb.text_frame, [[('IN ONE SENTENCE', 12, True, ITU_BLUE_DARK, False)]])
    tb = box(s, 1.1, 2.15, 11.1, 2.8)
    set_text(tb.text_frame, [[(text, 25 if len(text) <= 200 else 22, True, INK, False)]])
    practice_box(s, *practice, y=5.15)
    footer(s, tag, itu=True)
    notes(s, note)
    return s


def layer_stack(prs, head, layers, closing, tag, note):
    """Three or four layers drawn as a stack, top first; the last (bottom) layer dark."""
    s = add_slide(prs, LAYOUT_WHITE)
    title(s, head)
    n = len(layers)
    h = min(1.15, 4.3 / n)
    for i, (name, gloss) in enumerate(layers):
        last = i == n - 1
        node(s, 0.72, 1.65 + i * (h + 0.18), 11.9, h, name, [gloss],
             fill=ITU_BLUE_DARK if last else LIGHT, ink=WHITE if last else INK,
             head_ink=WHITE if last else ITU_BLUE_DARK, head_size=15, size=15)
    closing_line(s, closing, 1.65 + n * (h + 0.18) + 0.1)
    footer(s, tag)
    notes(s, note)
    return s


def paera_picture(prs, head, tag, note):
    """1.2 slide 2 — PAERA's picture: services on top, the four pillars, the foundation across
    the width. Plain text boxes only (the written guide carries figure F2)."""
    s = add_slide(prs, LAYOUT_WHITE)
    title(s, head)
    node(s, 0.72, 1.55, 11.9, 1.0, 'SERVICES',
         ['For example, registering a learner, paying a scholarship.'],
         fill=PANEL_GREY, head_size=15, size=14)
    w, gap, y, h = 2.75, 0.3, 2.8, 1.9
    pillars = [('ACCESS', 'Connections, devices, skills, and help for people who need it.'),
               ('DIGITAL DATA', 'The state registers.'),
               ('INTEROPERABILITY', 'The exchange of data between systems.'),
               ('DIGITAL IDENTITY', 'Who a person is.')]
    for i, (name, gloss) in enumerate(pillars):
        node(s, 0.72 + i * (w + gap), y, w, h, name, [gloss], fill=LIGHT, head_size=15,
             size=13)
    node(s, 0.72, 4.95, 11.9, 1.1, 'FOUNDATION',
         ['Governance and policy, with the legal framework.'],
         fill=ITU_BLUE_DARK, ink=WHITE, head_ink=WHITE, head_size=15, size=14)
    closing_line(s, 'Services stand on the pillars; the pillars stand on the foundation.', 6.25)
    footer(s, tag)
    notes(s, note)
    return s


def four_blocks(prs, head, tag, note):
    """1.10 slide 3 — one service, four blocks, on the data exchange layer. Text boxes only
    (the written guide carries figure F6)."""
    s = add_slide(prs, LAYOUT_WHITE)
    title(s, head)
    node(s, 0.72, 1.5, 11.9, 0.95, 'LEARNER REGISTRATION SERVICE',
         ['The first proof: a learner registered once.'], fill=PANEL_GREY, head_size=15, size=14)
    w, gap = 5.8, 0.3
    for i, (name, who) in enumerate([('REGISTRATION', 'Education\'s own — the sector builds it.'),
                                     ('LEARNER REGISTER', 'Education\'s own — the sector builds it.')]):
        node(s, 0.72 + i * (w + gap), 2.65, w, 1.1, name, [who], fill=LIGHT, head_size=15, size=14)
    for i, (name, who) in enumerate([('IDENTITY', 'The country\'s — used as it is.'),
                                     ('PAYMENTS', 'The country\'s — set up beside the service.')]):
        node(s, 0.72 + i * (w + gap), 3.95, w, 1.1, name, [who], fill=LIGHT, head_size=15, size=14)
    node(s, 0.72, 5.25, 11.9, 0.95, 'DATA EXCHANGE LAYER', ['Joins the four blocks to the service.'],
         fill=ITU_BLUE_DARK, ink=WHITE, head_ink=WHITE, head_size=15, size=14)
    foot_line(s, 'The written guide carries figure F6 for this picture.', 6.35)
    footer(s, tag)
    notes(s, note)
    return s


# ---------------------------------------------------------------- COVER (edit slide 1)
edit_cover(
    prs,
    title_text='Where your country stands,\nand what to build first',
    kicker='KP3 · Education DPI Roadmap · Module 1',
    blurb='Ten standalone videos for the person who commissions a DPI assessment and plans '
          'what to build first: what makes infrastructure foundational, the five domains on one '
          'page, the nine steps of the method, the desk review, the questionnaires, verification, '
          'scoring on five stages — and the order of building, down to the first proof for '
          'education.',
    length='~50 mins across 10 videos (1.1 – 1.10)',
    audience=AUDIENCE,
    panel_heading='WHAT THIS MODULE SETTLES',
    panel_items=['Four marks for anything called DPI',
                 'Five domains on one page',
                 'Nine steps, five answers each',
                 'Desk review · questionnaires · verification',
                 'Five stages, one table',
                 'Blocks, order, first proof'],
    panel_footer='4 marks · 5 domains · 9 steps · 5 stages · 1 first proof',
    note_text='Cover for the combined Module 1 deck. Each section that follows is one standalone '
              '~5 minute video. This is the Strategist-facing entry point to KP3 — the public-sector '
              'middle manager who commissions the assessment, defends its result to the minister '
              'and plans what to build first. Modules 2 to 5 then stand up the four blocks the '
              'first proof needs; Module 6 turns the proof into a national roadmap.')

# ---------------------------------------------------------------- AGENDA (edit slide 2)
edit_agenda(
    prs,
    header='Module 1 — ten videos',
    items=[
        ('1.1  What makes infrastructure foundational', '~5 min'),
        ('1.2  The five domains, as a map', '~5 min'),
        ('1.3  The roadmap method in nine steps: who does what', '~5 min'),
        ('1.4  Start at the desk: a first assessment from public sources, with AI', '~5 min'),
        ('1.5  Ask the people who run the systems: the five questionnaires', '~5 min'),
        ('1.6  Verify before you score', '~5 min'),
        ('1.7  Score each domain: five stages and one table', '~5 min'),
        ('1.8  Foundational blocks and the sector\'s own', '~5 min'),
        ('1.9  What comes first: the order of building', '~5 min'),
        ('1.10  The first proof for education: one service on four blocks', '~5 min'),
    ],
    message_paras=[
        'Test anything called DPI against four marks; see the country on one page of five '
        'domains; assess it in five steps a ministry can commission and defend.',
        'Then fund each block as what it is, build in PAERA\'s order, and prove one service '
        'first: a learner registered once, on four blocks.',
    ],
    note_text='Navigation slide for the combined deck; the videos ship standalone on YouTube. '
              '1.1 and 1.2 are the frame. 1.3 to 1.7 are the first five steps of the method — '
              'the plan, the desk review, the questionnaires, verification, scoring. 1.8 to 1.10 '
              'reason from the result to the first proof for education.')

delete_template_slides(prs, keep=2)


# ================================================================ 1.1
T = '1.1 · What makes infrastructure foundational'
MSG = ('Before you fund anything called "DPI", test it against four marks: it serves the whole '
       'society, many services can connect to it, it is built on open standards, and clear rules '
       'govern it.')
section('1.1', 'What makes infrastructure foundational', MSG,
        "VO: Donors and vendors all speak of DPI, digital public infrastructure. Some of what they "
        "offer is shared by the whole government. Some of it is one ministry's system under a new "
        "name. Before your ministry pays for either, you need a test you can defend.")

rows_block(prs, 'What DPI means',
           [('UNDP: a set of foundational digital systems', ''),
            ('The description from the G20 agreement of 2023',
             'Shared, secure and interoperable systems, built on open standards, serving a whole '
             'society, governed by legal frameworks and enabling rules.'),
            ('Three kinds so far',
             'Digital identity, digital payments, consent-based data sharing — and others '
             'emerging.')],
           None, T,
           "VO: UNDP defines digital public infrastructure as a set of foundational digital "
           "systems. Its compendium of 2023 gives the longer description from the G20 agreement of "
           "that year. These are shared digital systems, secure and interoperable. They can be "
           "built on open standards. They deliver services at the scale of a whole society. And "
           "legal frameworks and enabling rules govern them. The compendium names three kinds so "
           "far: digital identity, digital payments, and data sharing based on consent. It expects "
           "others to emerge.",
           numbered=False)

rows_block(prs, 'Four marks',
           [('It serves the whole society', 'Not limited by place or by group.'),
            ('Many services can connect to it', 'Many uses, many providers.'),
            ('It is built on open standards', 'Anyone may build on it.'),
            ('Clear rules govern it', 'To protect people and prevent misuse.')],
           'The rules decide who gains from it.',
           T,
           "VO: From that description the compendium draws four characteristics. Turn them into "
           "four marks you can test. First, it serves the whole society, not one region or one "
           "group. Second, many services can connect to it, from many providers. Third, it is "
           "built on open standards, so that anyone may build on it. Fourth, clear rules govern "
           "it, to protect people and prevent misuse. Do not treat the fourth mark as a "
           "formality. PAERA, the GovStack reference architecture, places governance, policy and "
           "law at the foundation of a country's digital infrastructure, under its four pillars. "
           "It also says that such infrastructure shapes what can be built on top of it. The "
           "rules decide who gains from it.\n\n"
           "Production cue: the core of the video — the test the listener carries away. Reveal "
           "the four rows one at a time. Hold it a beat longer.")

rows_block(prs, 'What shared infrastructure has done',
           [('India',
             'Adults with a transaction account rose from about one in four in 2008 to over 80 '
             'percent — a change estimated to have taken up to 47 years without DPI (World Bank '
             'report for the G20, 2023, page 2).'),
            ('Brazil',
             'By the end of February 2022, 15 months after launch, 67% of adults had made or '
             'received a payment through Pix (BIS Bulletin No 52, 2022, page 5).')],
           'Procurement rules can require interoperability, but they cannot deliver it. Only '
           'whole-of-government planning sees the total.',
           T,
           "VO: Why does the test matter? See what shared systems did. A World Bank report for the "
           "G20 looks at India. There, the share of adults with a transaction account rose from "
           "about one in four in 2008 to over 80 percent. The report estimates that without such "
           "infrastructure the same change could have taken up to 47 years. The report adds that "
           "other policies mattered as well. In Brazil, by the end of February 2022, 67 percent of "
           "adults had used Pix, the central bank's payment system, 15 months after its launch. "
           "Procurement rules can require interoperability, but they cannot deliver it. A project "
           "pays only for its own users. Only whole-of-government planning sees the total: the "
           "first ministry pays, and every ministry after it re-uses what was built.\n\n"
           "On screen: country names in plain typography only — no flags, emblems or logos.",
           numbered=False, bottom=4.9)

table_slide(prs, 'Progressa\'s systems against the four marks',
            ['System', 'Result', 'Reason'], [3.1, 2.9, 5.9],
            [('National identity (PNIA)', 'Foundational, with gaps',
              '78 per cent of adults covered; no identity number below the age of issue; no '
              'education service uses its sign-in.'),
             ('Payments', 'Not yet in place for government',
              'PayPro runs a fast-payment system; no government payments service is joined to Linkup.'),
             ('Linkup, the data exchange', 'Foundational, still a pilot',
              'Built for many bodies to connect; the ministry of education is not yet a member.'),
             ('Learner registry (PLR)', 'The sector\'s own, not yet authoritative',
              'Records for some learners only, served to the examination authority alone.'),
             ('PEMIS, the ministry\'s school information system', 'A sector application',
              'Holds totals, not learners; shares data by file export.')],
            T,
            "VO: Now apply the four marks to Progressa, the fictional country of these examples. "
            "Its national identity, run by PNIA, the identity authority, covers 78 per cent of "
            "adults. Children below the age of issue have no identity number, and no education "
            "service uses its sign-in yet. Foundational, with gaps. PayPro, the payment provider, runs a fast-payment system, but no payments service for government is joined to Linkup, the national data exchange, yet. So no education programme can pay through it. Linkup is built for many bodies to connect, "
            "but it is still a pilot, and the ministry of education is not a member. PLR, the "
            "learner registry, holds records for some learners only. It belongs to education and "
            "is not yet the authoritative register. PEMIS, the ministry's school information "
            "system, holds totals, not learners. It is a sector application.\n\n"
            "Production cue: the worked example — reveal the rows one at a time.",
            foot='Progressa is a fictional country. The screen is illustrative.')

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Test anything called DPI against four marks: the whole society, many services, open "
      "standards, clear rules. A system that misses a mark is not yet infrastructure, whatever it "
      "is called.",
      ('Screen your systems against the four marks of foundational infrastructure',
       'a table with one line of reasoning for each mark, for each system'))

sources_slide(prs, T, [
    'UNDP — Digital public infrastructure (web page)',
    'UNDP — Accelerating the SDGs through Digital Public Infrastructure: A Compendium (2023), '
    'pages 3 and 4',
    'PAERA v1.0 — sections 2.6 and 3.3.1',
    'World Bank — G20 Policy Recommendations for Advancing Financial Inclusion and Productivity '
    'Gains through Digital Public Infrastructure (2023), page 2',
    'BIS Bulletin No 52 (2022), page 5',
])


# ================================================================ 1.2
T = '1.2 · The five domains, as a map'
MSG = ('One page with five domains, a foundation of governance, policy and law carrying access, '
       'digital data, interoperability and digital identity, shows your minister what the '
       'country has and what it lacks.')
section('1.2', 'The five domains, as a map', MSG,
        "VO: Ask a ministry what the country has for digital government, and you often get a "
        "list of projects, each with its own donor. A list hides what is missing. One page with "
        "five domains shows your minister both.")

# The module's centrepiece: the one diagram of the module, in text boxes only.
paera_picture(prs, 'PAERA\'s picture of the infrastructure', T,
              "VO: PAERA, the GovStack reference architecture, gives the picture. It says that a country's digital governance infrastructure has two parts. The "
              "first is a foundation framework of governance, policy and legal components. The "
              "second is four pillars: Access, Digital Data, Interoperability and Digital "
              "Identity. Services, such as registering a learner or paying a scholarship, stand "
              "on the pillars. PAERA also explains why the pillars come early. They are the "
              "technical conditions that services need, and a country should address them "
              "before it pursues wide and ambitious plans.\n\n"
              "Production cue: the centrepiece of the module — the picture every later video "
              "scores against. Reveal the foundation first, then the four pillars, then the "
              "services on top. Hold it a beat longer. The written guide carries figure F2.")

rows_block(prs, 'Five domains, with their codes',
           [('GOV', 'Governance and policy, with the legal framework.'),
            ('ACC', 'Access.'),
            ('DAT', 'Digital Data.'),
            ('INT', 'Interoperability.'),
            ('IDN', 'Digital Identity.')],
           'Reading the foundation as a fifth domain is this course\'s own reading of PAERA.',
           T,
           "VO: The method reads the foundation as a fifth domain beside the four pillars. That "
           "reading is this course's own. Its reason is practical: a pillar without a law, a budget "
           "and a body that decides is not yet infrastructure. So there are five domains, each "
           "with a short code. GOV is governance and policy, with the legal framework. ACC is "
           "access: connections, devices, skills, and help for people who need it. DAT is digital "
           "data, with the state registers. INT is interoperability, the exchange of data between "
           "systems. IDN is digital identity.",
           numbered=False, bottom=5.9)

table_slide(prs, 'Progressa\'s frame page',
            ['Domain', 'Who answers for Progressa'], [2.2, 9.7],
            [('GOV', 'The ministry\'s planning directorate, with PDGA for the national rules.'),
             ('ACC', 'The ministry\'s ICT unit, with the district education offices.'),
             ('DAT', 'The ministry\'s statistics unit, which runs PEMIS, with PNEA, the '
                     'examination authority.'),
             ('INT', 'PDGA, which operates Linkup, with the ministry\'s ICT unit.'),
             ('IDN', 'PNIA, the national identity authority, with the ministry\'s ICT unit.')],
            T,
            "VO: Here is the map for Progressa, the fictional country of these examples. It is step one of the method, and the first output of the assessment: one page that fixes its scope. For each domain "
            "it names the body that answers. For digital data, that is the statistics unit of "
            "the ministry of education, which runs PEMIS, its school information system, with "
            "PNEA, the examination authority. For interoperability, it is PDGA, the digital "
            "government authority, which operates Linkup, the national data exchange. The page "
            "also records what is there and what is not. Linkup runs as a pilot. PLR, the learner "
            "registry, holds records for some learners only, and is not yet the authoritative "
            "learner register.",
            foot='Progressa is a fictional country. Worked example E1.', size=14)

panels(prs, 'One page both sides can read',
       ('THE MINISTER READS', ['The gaps in law, money and ownership.']),
       ('THE ARCHITECT READS', ['The gaps in systems and data.']),
       'When they decide what comes first, both point to the same part of the same page.',
       T,
       "VO: Why one page? PAERA describes enterprise architecture as documents that "
       "look at an organisation from business and IT together. Their purpose is to close the gap "
       "in communication between the two. The five-domain map does the same for a country's "
       "infrastructure. It gives the policy side and the technical side one shared language. The "
       "minister reads the gaps in law, money and ownership. The architect reads the gaps in "
       "systems and data. When they decide what comes first, both point to the same part of the "
       "same page. A list of projects cannot do this, because it shows only what someone chose to "
       "fund.\n\n"
       "Production cue: the pivotal slide of this video — the shared-language argument. Hold it "
       "a beat longer.",
       height=2.4)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Five domains on one page: a foundation of governance, policy and law, and four pillars "
      "on it. It shows the minister what exists and what is missing.",
      ('Place your systems, laws and public bodies on the five-domain map',
       'a map of the five domains with every item placed and its reasoning'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 2.3, 2.5, 3.1, 3.3.1 and 3.4.1 to 3.4.4',
])


# ================================================================ 1.3
T = '1.3 · The roadmap method in nine steps: who does what'
MSG = ('A roadmap is produced in nine steps, and for each step you can say who acts, what goes '
       'in, what comes out, who decides and how the result is checked.')
section('1.3', 'The roadmap method in nine steps: who does what', MSG,
        "VO: A minister who commissions a roadmap will ask three things. Who does the work? When "
        "must I decide? How will anyone know the result is right? Nine named steps let you answer "
        "all three before the work starts.")

rows_block(prs, 'Five questions for every step',
           [('Who acts?', ''), ('What goes in?', ''), ('What comes out?', ''),
            ('Who decides?', ''), ('How is the result checked?', '')],
           'And which tool or template the step uses.',
           T,
           "VO: For each step you can say five things: who acts, what goes in, what comes out, "
           "who decides, and how the result is checked. Add one more: the tool or template the "
           "step uses. The roles are few and they recur. The facilitator plans and runs the "
           "assessment and keeps its evidence. A respondent team answers for one domain, with a "
           "section lead for each part of its questionnaire. Focal points in other bodies confirm "
           "what is said about their systems. Reviewers check each step before the next one "
           "starts. At the end, an adopting authority makes the roadmap the government's own.",
           bottom=5.9)

table_slide(prs, 'Steps 1 to 5: where the country stands',
            ['Step', 'What comes out', 'Tool or template'], [2.7, 5.0, 4.2],
            [('1. Frame', 'One page: the five domains and who answers for each', 'The frame page'),
             ('2. Automatic assessment', 'First findings and gaps, each with its source, unverified',
              'The question bank and the workflow that runs it'),
             ('3. Questionnaires', 'One answered questionnaire per domain, with its evidence',
              'The five questionnaires and their two guides'),
             ('4. Verification', 'A verified position on every point',
              'The templates for verification'),
             ('5. Scoring', 'Five domain reports and the maturity table',
              'The scoring criteria and the scoring prompt')],
            T,
            "VO: The first five steps find out where the country stands. Step one frames the work "
            "on one page. In Progressa, the fictional country of these examples, the ministry of "
            "education sponsors it together with PDGA, the digital government authority. Step two "
            "reads what the country has published, with an AI assistant; nothing it drafts is yet "
            "a finding. Step three sends one questionnaire to each domain. Step four checks every "
            "answer before anything is scored; in Progressa, four verification sessions took "
            "place in the eighth week. Step five scores each domain on five stages and puts the "
            "result on one table. An assistant may propose each stage; a person decides on the "
            "evidence, and a reviewer checks.\n\n"
            "Production cue: the pivotal slide of this video — the method table, first half. The "
            "written guide carries figure F3 and the full table. Hold it a beat longer.")

table_slide(prs, 'Steps 6 to 9: from the result to an adopted roadmap',
            ['Step', 'What comes out', 'Who decides'], [2.9, 4.3, 4.7],
            [('6. Gaps and priorities', 'The gap register, ranked',
              'A workshop of the bodies on the frame page'),
             ('7. Roadmap', 'The roadmap over time, with its dependencies',
              'Dates stay proposals until the budget and the owners of the work confirm them'),
             ('8. Investment breakdown', 'The investment case in four sheets',
              'What is funded first, cost against reuse'),
             ('9. Validation and revision', 'Every comment answered; the revised documents',
              'The adopting authority adopts the roadmap')],
            T,
            "VO: The last four steps turn the result into a roadmap. Step six turns findings into "
            "gaps and ranks them, and a workshop of the bodies concerned decides the order. Step "
            "seven writes the roadmap over time. Its dates stay proposals until the budget and "
            "the owners of the work confirm them. Step eight costs the roadmap in four sheets "
            "that a finance ministry can read. Step nine sends everything back to the bodies for "
            "comment, answers every comment in writing, and revises. In Progressa, the ministry of education and PDGA adopt the roadmap, with the cabinet where the budget requires it. Module 6 takes these four steps in full.")

rows_block(prs, 'Where the nine steps come from',
           [('PAERA 5.3', 'The nine questions a national assessment must answer; no scale.'),
            ('PAERA 5.4', 'Eight steps for assessing one organisation, from criteria to '
                          'continuous improvement.'),
            ('UNDP Playbook, page 23', 'Six steps from national priorities to a roadmap.')],
           'The sequence of nine steps is this course\'s own.',
           T,
           "VO: Where do the nine steps come from? No public source gives them in this form. "
           "PAERA lists nine questions that a national assessment must answer, and gives no scale. It gives eight steps for assessing one organisation, from setting the criteria to continuous improvement. It also lists seven signs of low maturity, such as a lack of digital data. UNDP's playbook gives six "
           "steps that lead from national priorities to a roadmap. The sequence of nine steps is "
           "this course's own, drawn from its authors' implementation experience in several countries.\n\n"
           "Production cue: marked place 1 of 2 (naming) — if ITU agrees to name the country in "
           "which the method was first applied, one sentence of provenance is added to the end "
           "of this narration in the bundle; nothing else in the video changes.",
           numbered=False)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Nine steps, and for each one five answers: who acts, what goes in, what comes out, "
      "who decides, and how the result is checked. A minister can follow that.",
      ('Turn your list of institutions and a start date into an assessment plan',
       'a stakeholder map, a schedule by step, a plan of sessions and the list of documents to '
       'request'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 3.1.3, 5.3 and 5.4',
    'UNDP — The DPI Approach: A Playbook (2023), page 23',
])


# ================================================================ 1.4
T = '1.4 · Start at the desk: a first assessment from public sources, with AI'
MSG = ('An AI-assisted review of what your country has already published gives you first '
       'findings and gaps before you ask anyone a question, each tied to its source, and it is a '
       'draft to verify, not a verdict.')
section('1.4', 'Start at the desk: a first assessment from public sources, with AI', MSG,
        "VO: An assessment usually starts by asking officials for documents, and then waiting. "
        "Start instead with what your country has already published. In a few days an AI "
        "assistant gives you a first picture, and a list of points to check.")

rows_block(prs, 'The question bank',
           [('62 questions', 'Arranged by the five domains and their 26 sub-components.'),
            ('Each question', 'Its code, the sub-component it informs, the evidence that answers '
                              'it, what a good answer looks like.'),
            ('Example — DAT-Q03', 'Is there an authoritative register of learners, and which body '
                                  'keeps it?')],
           None, T,
           "VO: The desk review runs on a question bank. The bank in the assessment toolkit holds "
           "62 questions, arranged by the five domains and their 26 parts, called sub-components. "
           "Each question has a code. It names the part it informs, the documents that answer it, "
           "and what a good answer looks like. Take question DAT-Q03: is there an authoritative "
           "register of learners, and which body keeps it? A good answer is one register, kept by "
           "a named body under law, linked to the national identity. PAERA, the GovStack "
           "reference architecture, lists the subjects such an assessment must cover. It also says that GovStack has a tool for a quick assessment, which it does "
           "not name.",
           numbered=False)

rows_block(prs, 'The workflow, in five stages',
           [('Collect practice', 'Evidence notes, each with its source and date.'),
            ('Draft findings', 'One draft result per question, with a provisional stage and a '
                               'confidence between 0 and 1.'),
            ('Draft the gap analysis', 'The desk gaps.'),
            ('Make the report', 'One chapter per domain.'),
            ('Clean and sort', 'The sorted list of desk gaps.')],
           'A person reads the output of each stage before the next stage runs.',
           T,
           "VO: The assistant works in five stages. It collects what the public sources say, each "
           "note with its source and date. It drafts one result per question, with a provisional "
           "stage and a confidence between zero and one. It writes the gap between each result "
           "and good practice. It assembles a desk report, one chapter per domain. Then it "
           "removes duplicate gaps and sorts them, so that each can go to the right person. A "
           "person reads the output of every stage before the next one runs. The assistant "
           "drafts; it does not decide.\n\n"
           "Production cue: the pivotal slide of this video. Reveal the stages one at a time; let "
           "the closing line land.",
           bottom=5.9)

rows_block(prs, 'Progressa\'s desk run',
           [('AF-DAT-01', 'PLR holds enrolment records for some learners; enrolment is counted '
                          'from school returns. Basic; confidence 0.8.'),
            ('AF-INT-01', 'Linkup runs as a pilot; PNEA, PLR and PNIA are among its members; '
                          'MoEYS is not. Systematic; confidence 0.7.'),
            ('AF-DAT-02', 'The yearbook appears to report learner-level data for secondary '
                          'schools. Systematic; confidence 0.4.'),
            ('Desk gaps', 'DG-03: PLR not yet the authoritative learner register. DG-05: MoEYS '
                          'not a member of Linkup.')],
           'Every desk result is unverified. Worked example E2, illustrative.',
           T,
           "VO: Here is the run on the public record of Progressa, the fictional country of these "
           "examples. It read six sources, from the education statistics yearbook to the data "
           "protection act. Desk result AF-DAT-01 says that PLR, the learner registry, holds "
           "enrolment records for some learners. Enrolment is still counted from school returns. "
           "Provisional stage Basic, confidence 0.8. AF-INT-01 says that Linkup, the national data "
           "exchange, runs as a pilot, and that MoEYS, the ministry of education, is not a member: "
           "confidence 0.7. AF-DAT-02 suggests learner-level data for secondary schools, with a "
           "confidence of only 0.4. The run ends with desk gaps, such as DG-03: PLR not yet the "
           "authoritative learner register. It also reads the record against PAERA's seven signs of low maturity. Lack of digital data: yes, for learners.",
           numbered=False)

storyboard_slide(prs, 'Demonstration: the desk assessment on Progressa',
                 ['The settings of the run — Progressa, education, the six public sources with '
                  'their dates, the digital data questions loaded from the bank',
                  'Stage 1, collect practice — the evidence notes, each with its source and date',
                  'Stage 2, draft findings — AF-DAT-01 (Basic, confidence 0.8) and AF-DAT-02 '
                  '(Systematic, confidence 0.4), both marked unverified',
                  'Stages 3 and 4 — the desk gaps DG-03 and DG-04, and the report\'s chapter on '
                  'digital data',
                  'Stage 5, clean and sort — the sorted desk gaps, each with the unit it is put '
                  'to; AF-DAT-02 marked to check first'],
                 T,
                 "VO: The demonstration follows one domain, digital data, through the five "
                 "stages, from the question bank to the sorted desk gaps. A good run passes when every draft result names its question and its "
                 "source, carries a provisional stage and a confidence, and is marked unverified. "
                 "Nothing in it counts as a finding. The point with the lowest confidence goes to "
                 "the officials first. In Progressa, that point, AF-DAT-02, turned out to be "
                 "wrong when the officials showed their system.\n\n" + STORYBOARD_NOTE)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Read what the country has published, with AI, before you ask anyone. You get first "
      "findings and gaps, each tied to its source: a draft to verify, not a verdict.",
      ('Run a desk assessment from the question bank and your country\'s public sources',
       'a table of draft findings, each with its source, a provisional stage, a confidence and '
       'the mark UNVERIFIED'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 3.1.3 and 5.3',
    'The question bank and the workflow that runs it — this course\'s own instruments, in the '
    'assessment toolkit',
])


# ================================================================ 1.5
T = '1.5 · Ask the people who run the systems: the five questionnaires'
MSG = ('Five questionnaires, one for each domain, put the same questions to the people who run '
       'the systems and ask for evidence with every answer.')
section('1.5', 'Ask the people who run the systems: the five questionnaires', MSG,
        "VO: The public record tells you what a country says about its systems. The people who "
        "run them know what the systems actually do. Ask them, in writing, and ask for the "
        "document behind every answer.")

rows_block(prs, 'One questionnaire for each domain',
           [('Q-GOV', 'Governance and policy, with the legal framework.'),
            ('Q-ACC', 'Access.'),
            ('Q-DAT', 'Digital data.'),
            ('Q-INT', 'Interoperability.'),
            ('Q-IDN', 'Digital identity.')],
           '93 questions in all, grouped by the same 26 sub-components as the question bank.',
           T,
           "VO: There are five questionnaires, one for each domain. PAERA, the GovStack reference "
           "architecture, lists nine questions that a national assessment must answer. Each questionnaire expands one or more of them. The digital data "
           "questionnaire, for example, expands PAERA's question on the status of the national "
           "infrastructure for managing digital data. Together the five ask 93 questions. They "
           "cover the same 26 parts, or sub-components, as the desk review of the public record, "
           "so every answer can be set beside the desk result for the same part. PAERA names questionnaires, interviews and document review as tools of an assessment.",
           numbered=False, bottom=5.9)

rows_block(prs, 'How a question asks for evidence',
           [('D2.2  Is there one authoritative record of each learner?', ''),
            ('Describe', 'Where a learner\'s identity and enrolment are recorded today.'),
            ('Indicate', 'Whether the record is on paper, in a spreadsheet or in a system.'),
            ('Provide', 'A blank copy of the form or screen used.')],
           'An answer with a document behind it counts for more. And "we do not have this" is a '
           'useful answer.',
           T,
           "VO: Each question opens with one short lead question, and then asks for three things: "
           "what to describe, what to indicate, and what to provide. Take question D2.2 of the "
           "digital data questionnaire: is there one authoritative record of each learner? "
           "Describe where a learner's identity and enrolment are recorded today. Indicate "
           "whether the record is on paper, in a spreadsheet or in a system. Provide a blank copy "
           "of the form or screen used. An answer with a document behind it counts for more than "
           "an answer without one. And 'we do not have this' is a useful answer.\n\n"
           "Production cue: the pivotal slide of this video — the team's blank instrument, as the "
           "toolkit words it. Hold it a beat longer.",
           numbered=False)

panels(prs, 'Two guides with every questionnaire',
       ('FOR THE RESPONDENT',
        ['Who answers, what to prepare, how to answer, the five stages for describing, the checks '
         'before submitting.']),
       ('FOR THE FACILITATOR',
        ['Before the session, how to run it, what to look for, how to record answers and '
         'evidence, difficult situations, checks before scoring.']),
       'Respondents describe; they do not score.',
       T,
       "VO: Every questionnaire comes with two guides. The guide for the respondent says who "
       "answers, which documents to collect, how to answer, and what to check before submitting. "
       "In a typical case the work takes about 13 working days, and the head of the responding "
       "unit signs the answers. The guide for the facilitator says how to prepare, how to run a "
       "session, and how to record each answer with its evidence. Respondents describe; they do "
       "not score. Each questionnaire also carries the points from the desk review, so that the "
       "respondents can confirm or correct them.",
       height=2.8)

rows_block(prs, 'Progressa\'s answer on learner records',
           [('There is no single record of a learner.', ''),
            ('PLR holds enrolment records for some learners and serves them to PNEA alone',
             'Primary schools keep class registers on paper.'),
            ('Schools send totals by grade and sex once a year',
             'PEMIS holds no record of individual learners.')],
           'Documents: the blank annual school return form; the PEMIS data dictionary. Worked '
           'example E3, illustrative.',
           T,
           "VO: Here is how Progressa, the fictional country of these examples, answered that "
           "question. The statistics unit of its ministry of education replied, with PNEA, the "
           "examination authority, in week six of the assessment. There is no single record of a "
           "learner. PLR, the learner registry, holds enrolment records for some learners and "
           "serves them to PNEA alone. Primary schools keep class registers on paper. Each school "
           "sends totals by grade and sex once a year, and PEMIS, the ministry's school "
           "information system, holds no record of individual learners. The answer names its "
           "documents: the blank annual school return form and the PEMIS data dictionary.",
           numbered=False)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Five questionnaires, one for each domain, put the questions to the people who run the "
      "systems. Every answer names its evidence, or says plainly that none exists.",
      ('Read a filled questionnaire against its evidence',
       'a table with one row per question — what the evidence supports, what is stated without '
       'evidence, the question to ask next'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 5.3 and 5.4',
    'The questionnaires and their guides — this course\'s own instruments, in the assessment toolkit',
])


# ================================================================ 1.6
T = '1.6 · Verify before you score'
MSG = ('An answer becomes a finding only when it has been checked against documents, in '
       'interviews and workshops, and in a validation session with the people who gave it.')
section('1.6', 'Verify before you score', MSG,
        "VO: A questionnaire answer is what a body says about itself. A desk result is what the "
        "public record suggests. Neither is a finding yet. Score them as they stand, and the "
        "first official who disagrees can overturn your result.")

rows_block(prs, 'Six activities, in order',
           [('Receive the signed questionnaires', ''),
            ('Review every document they name', 'Against the desk results.'),
            ('Clarification interviews', 'On each disputed point.'),
            ('Technical workshops', 'See the systems working.'),
            ('Synthesis', 'A verified position for every point.'),
            ('Final validation', 'With every body.')],
           None, T,
           "VO: Verification runs in six activities, in order. The facilitator receives the "
           "signed questionnaires. The facilitator then reads every document they name, and "
           "compares it with the desk results. Each disputed point goes to the person who owns "
           "it, in a clarification interview. In technical workshops the team sees the systems "
           "working: a screen, a report, a log. The facilitator then writes a verified position "
           "for every point. Last, a validation workshop with all the bodies closes what is still "
           "disputed. PAERA, the GovStack reference architecture, names the same tools: document review, interviews and focus groups, followed by the conduct of "
           "the assessment and the analysis of its results.",
           bottom=6.6)

rows_block(prs, 'The evidence ladder',
           [('A system seen working', ''),
            ('An official document in force', ''),
            ('Official statistics', ''),
            ('An internal report', ''),
            ('A statement in an interview', '')],
           'When sources disagree, the higher level wins. A verified position needs one source at '
           'level 1, or two independent sources at levels 2 or 3.',
           T,
           "VO: When sources disagree, a rule settles it: the evidence ladder. A system seen "
           "working stands highest. Then comes an official document in force, such as a law or an "
           "approved budget. Then official statistics, then an internal report, and last a "
           "statement in an interview. The higher level wins. A position counts as verified only "
           "with one source at the first level, or with two independent sources at the second or "
           "third. Where two sources at the same level still disagree, the point is marked "
           "disputed and goes to the validation workshop.\n\n"
           "Production cue: the pivotal slide of this video — the rule that settles "
           "disagreements. Hold it a beat longer.",
           bottom=5.7)

rows_block(prs, 'Four templates',
           [('The map of the bodies involved', 'Who takes part, and how.'),
            ('The session plan', 'Each interview and workshop, the points to clarify, the '
                                 'evidence to bring.'),
            ('The reconciled-response table', 'The verified position on every point, with its '
                                              'evidence and its status.'),
            ('The agenda of the final validation workshop', 'Half a day, chaired by the '
                                                            'validation chair.')],
           None, T,
           "VO: Four blank templates carry the work. The map of the bodies involved says who "
           "takes part, and how. The session plan sets out each interview and workshop, with the "
           "points to clarify and the evidence to bring. The reconciled-response table holds the "
           "verified position on every point, with its evidence and its status: verified, "
           "corrected or disputed. The agenda of the final validation workshop orders the half "
           "day in which every body sees the result. A validation chair, a senior official named "
           "by the sponsor of the assessment, decides the points still disputed.",
           numbered=False)

rows_block(prs, 'Three points Progressa\'s team verified',
           [('AF-DAT-02 — corrected', 'PEMIS holds totals by district, grade and sex; PEMIS '
                                      'shown in session V-01 (level 1).'),
            ('AF-INT-01 — verified', 'MoEYS is not a member of Linkup and runs no service on it; '
                                     'member list exported by PDGA in session V-02 (level 1).'),
            ('AF-IDN-02 — verified', 'Only PNEA reads persons from PNIA, through the read by '
                                     'national number; no education service is a client of '
                                     'PNIA\'s sign-in; shown in session V-03 (level 1).')],
           'Worked example E4, illustrative.',
           T,
           "VO: Here is how it worked in Progressa, the fictional country of these examples. Sessions were held in week eight. Here are three of the points they settled. The desk review had suggested that PEMIS, the "
           "ministry's school information system, held learner-level data. In session V-01 the "
           "team saw PEMIS working: it holds totals by district, grade and sex. The point was "
           "corrected. In session V-02, PDGA, which operates Linkup, the data exchange, exported "
           "its member list: MoEYS, the ministry of education, is not a member. Verified. In "
           "session V-03, PNIA, the identity authority, showed its service log. Only PNEA, the "
           "examination authority, reads persons from PNIA, by national number. No education "
           "service uses PNIA's sign-in. Verified. All three were settled at the first level: a "
           "system seen working.",
           numbered=False)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: An answer becomes a finding only after it is checked: in documents, in interviews and "
      "workshops, and in a validation session with the people who gave it.",
      ('Set two disagreeing sources side by side and draft the clarification question',
       'a two-column comparison of the sources with their levels, the evidence that would settle '
       'the point, and the question for the clarification interview'))

sources_slide(prs, T, [
    'PAERA v1.0 — section 5.4, steps 2 to 4',
    'The verification templates and the evidence ladder — this course\'s own instruments, in the '
    'assessment toolkit',
])


# ================================================================ 1.7
T = '1.7 · Score each domain: five stages and one table'
MSG = ('Score each domain on five published stages against written criteria, so that two '
       'assessors reach the same result and the whole picture fits on one table.')
section('1.7', 'Score each domain: five stages and one table', MSG,
        "VO: Two assessors who read the same evidence should reach the same score. If they do "
        "not, the score is an opinion, and a minister can set it aside. Written criteria on a "
        "published scale prevent that.")

rows_block(prs, 'Five published stages',
           [('Basic', '0 to 1.'), ('Opportunistic', '1 to 2.'), ('Systematic', '2 to 3.'),
            ('Differentiating', '3 to 4.'), ('Transformational', '4 to 5.')],
           'UNDP Digital Development Compass: all scores rescaled between 0 and 5, with no zero '
           'scores.',
           T,
           "VO: The scale is the one UNDP publishes in its Digital Development Compass. It has "
           "five stages, under these names: Basic, Opportunistic, Systematic, Differentiating and "
           "Transformational. The Compass rescales all scores between zero and five, with no zero "
           "scores, so that each stage takes one unit of that range. The domains you score are those of PAERA, the GovStack reference architecture, which gives no scale of its own for a national assessment.",
           numbered=False, bottom=5.7)

rows_block(prs, 'The scoring rule',
           [('Each sub-component has written criteria', 'At Basic, Systematic and '
                                                        'Transformational.'),
            ('Each sub-component is placed at one stage', 'From verified positions only.'),
            ('Each stage scores the middle of its unit', '0.5, 1.5, 2.5, 3.5, 4.5.'),
            ('A domain\'s score is the mean of its sub-components', 'Its stage is the unit the '
                                                                    'score falls in.')],
           None, T,
           "VO: The rule that lays the Compass over the five domains is this course's own, and it "
           "has four parts. Each of the 26 sub-components has written criteria at three levels: "
           "Basic, Systematic and Transformational. Meeting part of the Systematic criterion "
           "places a sub-component at Opportunistic; part of the Transformational one, at "
           "Differentiating. Only verified positions are scored. Each stage scores the middle of "
           "its unit, from 0.5 for Basic to 4.5 for Transformational. A domain's score is the mean "
           "of its sub-components, to one decimal place, and its stage is the unit the score "
           "falls in.")

table_slide(prs, 'Progressa\'s maturity table',
            ['Domain', 'Score', 'Stage'], [4.2, 2.6, 5.1],
            [('GOV', '1.9', 'Opportunistic'), ('ACC', '1.3', 'Opportunistic'),
             ('DAT', '1.3', 'Opportunistic'), ('INT', '2.1', 'Systematic'),
             ('IDN', '2.3', 'Systematic')],
            T,
            "VO: Here is the result for Progressa, the fictional country of these examples, on "
            "one table. Governance, 1.9, Opportunistic. Access, 1.3, Opportunistic. Digital data, "
            "1.3, Opportunistic. Interoperability, 2.1, Systematic. Digital identity, 2.3, "
            "Systematic. Each row also names a main strength and a development priority. For "
            "digital data, the strength is a working national identity register, and the "
            "priority is to make PLR, the learner registry, the authoritative learner register. "
            "The digital data score is the mean of six sub-components. Five sit at Opportunistic "
            "and one, data quality and master data, at Basic. That gives 1.3.\n\n"
            "Production cue: the pivotal slide of this video — the whole picture on one table. "
            "The written guide carries figure F4. Reveal the rows one at a time.",
            foot='Illustrative. Worked example E5.', size=15)

storyboard_slide(prs, 'Demonstration: scoring Progressa\'s digital data domain',
                 ['The two inputs — the criteria of digital data, D1 to D6 with five rows each, '
                  'and the domain\'s verified answers with their evidence',
                  'Part A — one row per sub-component: proposed stage, score, and the quotations '
                  'for the parts of the criterion met and not met',
                  'Part B, the arithmetic — (1.5 + 1.5 + 0.5 + 1.5 + 1.5 + 1.5) ÷ 6 = 1.3, '
                  'Opportunistic',
                  'Parts C and D — the answers not used, with the reason, and any two pieces of '
                  'evidence that disagree, quoted side by side',
                  'The record of confirmation — the stage proposed and the stage confirmed for '
                  'each sub-component, with the reason for any change'],
                 T,
                 "VO: The demonstration applies the scoring prompt to Progressa's digital data "
                 "domain. The prompt proposes a stage for each "
                 "sub-component, and quotes the evidence for each criterion. The assessor "
                 "confirms or changes each stage in a written record. A good run passes when every stage rests on quoted, verified evidence and the arithmetic gives 1.3. One correction "
                 "shows why the rule matters. Progressa's draft placed interoperability at 2.3, "
                 "from an unverified desk result. But a session had already verified that the "
                 "ministry of education is not a member of Linkup, the data exchange. Once that "
                 "was pointed out, the score fell to 2.1. Its stage stayed Systematic.\n\n"
                 + STORYBOARD_NOTE)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Score every domain on five published stages against written criteria. Two assessors "
      "then reach the same result, and the whole picture fits on one table.",
      ('Propose a stage for each sub-component with the scoring prompt',
       'a proposed stage and score for every sub-component, each criterion quoted against the '
       'evidence'))

sources_slide(prs, T, [
    'UNDP — Digital Development Compass: Methodology (web page)',
    'PAERA v1.0 — sections 3.3.1, 5.1, 5.3 and 5.4',
    'The criteria and the scoring rule — this course\'s own, in the assessment toolkit',
])


# ================================================================ 1.8
T = '1.8 · Foundational blocks and the sector\'s own'
MSG = ('Identity, payments and data exchange are foundations for every sector, while a learner '
       'register and its registration service belong to education and stand on them, and each '
       'must be funded as what it is.')
section('1.8', 'Foundational blocks and the sector\'s own', MSG,
        "VO: Once the country is scored, the money requests come. The first question for each is whose system it is. Is it shared by every sector, or does it belong to one? Get that wrong, and the "
        "budget pays for the same thing twice.")

layer_stack(prs, 'Three layers',
            [('SECTOR APPLICATIONS', 'Digital education, digital health, and others.'),
             ('CORE DPI', 'Digital identity, digital payments, consent-based data sharing — '
                          'others emerging.'),
             ('GOVERNANCE FOUNDATIONS', 'Leadership, institutions, policy, law, engagement, '
                                        'technical expertise.')],
            'Education\'s own systems stand on the core. They are not part of it.',
            T,
            "VO: UNDP's compendium on digital public infrastructure draws three layers. At the "
            "bottom are the governance foundations: leadership, accountable institutions, policy, "
            "law, engagement with users, and technical expertise. In the middle is the core of "
            "the infrastructure: digital identity, digital payments and data sharing based on "
            "consent, with others emerging. On top are the sector applications, and digital "
            "education is one of them, beside digital health and others. So in UNDP's picture, "
            "education's own systems stand on the core. They are not part of it.\n\n"
            "Production cue: reveal the layers bottom to top.")

rows_block(prs, 'Where the registers sit',
           [('Basic national infrastructure', 'Digital ID, payments, legal data registries.'),
            ('A state registry needs two building blocks', 'Registration and Digital Registry.'),
            ('The main state registries include an Education Register',
             'One of the registers a country should have.')],
           'A learner register is a state register — but it belongs to education, and so does the '
           'registration service that writes to it.',
           T,
           "VO: The registers come next. Digital ID, payments and legal data registries count as basic national "
           "infrastructure. A country should have a set of main state registries, "
           "and an Education Register is one of them. Digitalising a state registry needs two building blocks: Registration and a Digital Registry. So a learner "
           "register is a state register, but it belongs to education, and so does the "
           "registration service that writes to it.",
           numbered=False)

block(prs, 'Identity is foundational; a learner\'s identity is not',
      ['The Identity block covers foundational identity: proof of identity for a wide variety of '
       'public and private services.',
       'A functional identity, for one purpose or one sector such as education, is outside its '
       'scope.'],
      'Education does not build an identity of its own. Its services use the national one.',
      T,
      "VO: Identity shows why the line matters. The GovStack Identity specification covers "
      "foundational identity: proof of who a person is, for a wide variety of public and private "
      "services. It places functional identity, for one purpose or one sector, outside its "
      "scope, and it names education as one such sector. It also notes that where a proper "
      "foundational identity exists, a separate functional one is no longer needed. So education "
      "does not build an identity of its own. Its services use the national one.")

table_slide(prs, 'Progressa, sorted into the three layers',
            ['Layer', 'Progressa', 'Funded as'], [2.5, 5.9, 3.5],
            [('Governance foundations', 'PDGA\'s mandate for the national exchange; the education '
                                        'act; the data protection act',
              'Law and budget, decided above any one project'),
             ('Core', 'PNIA\'s identity and sign-in; Linkup, the data exchange; the Payments '
                      'block, which reaches PayPro through its payer bank',
              'Once, for every sector'),
             ('Education', 'PLR, the learner register; the learner registration service; PEMIS',
              'By the sector, standing on the core')],
            T,
            "VO: Now sort Progressa, the fictional country of these examples. Its governance "
            "foundations include the mandate of PDGA, the digital government authority, for the "
            "national exchange, the education act and the data protection act. Its core is the "
            "identity and sign-in of PNIA, the identity authority, and Linkup, the data exchange. "
            "A Payments block belongs there too: it reaches PayPro, the payment provider, through "
            "its payer bank, and it is not yet on Linkup. Education's own layer holds PLR, the "
            "learner register, and PEMIS, the school information system. The learner "
            "registration service, still to be built, belongs there too. Fund the core once, for "
            "every sector. Fund the register and the service as education's, built on the core. "
            "A sector's register funded as national infrastructure takes money from the core. A "
            "core funded inside one project leaves every other sector to build it again.\n\n"
            "Production cue: the pivotal slide of this video — the funding rule made concrete. "
            "Hold it a beat longer.",
            foot='Illustrative.', size=14)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Identity, payments and data exchange serve every sector. The learner register and its "
      "registration service are education's own, standing on them. Fund each as what it is.",
      ('Decide whether a proposed system is a shared block, a sector\'s register or a sector\'s '
       'application',
       'a three-question assessment with a recommended answer'))

sources_slide(prs, T, [
    'UNDP — Accelerating the SDGs through Digital Public Infrastructure: A Compendium (2023), '
    'page 4, Exhibit 2',
    'PAERA v1.0 — section 3.3.3, Annex 1 (A1.2.5) and Annex 3',
    'GovStack Identity Building Block specification, Version 2.0 (December 2025), sections 2 '
    'and 3',
])


# ================================================================ 1.9
T = '1.9 · What comes first: the order of building'
MSG = ('Build identity and payments first, bring in data exchange and registration with the '
       'first priority services, digitalise the registers next, and show one visible result '
       'early so that support holds.')
section('1.9', 'What comes first: the order of building', MSG,
        "VO: Money for digital government comes in pieces: a donor here, a budget line there. "
        "Each piece is spent on something. Unless someone has fixed the order of building, the "
        "pieces arrive in an order that nobody chose.")

rows_block(prs, 'PAERA\'s four phases',
           [('Inception', 'Identity, Payment, and no-code or low-code development.'),
            ('High-priority use cases', 'The Information Mediator, Registration and other '
                                        'infrastructural blocks, with the first services.'),
            ('Initial transformation', 'All main state registries digitalised, in no more than 2 '
                                       'to 3 years.'),
            ('Mass-scale transformation', 'Every government service.')],
           'Each phase puts in place the blocks that the next one needs.',
           T,
           "VO: PAERA, the GovStack reference architecture, gives an order in four phases. In the "
           "first phase, inception, the country puts in place the Identity and Payment building "
           "blocks. It adds a platform for building services with little or no code. In the "
           "second, it picks high-priority services and builds them fast. They are built on the "
           "Information Mediator, which is the data exchange layer, on Registration and on other "
           "shared blocks. In the third, all main state registries are digitalised, in no more "
           "than two to three years. The fourth completes the digitalisation of every government "
           "service. Each phase puts in place the blocks that the next one needs.\n\n"
           "Production cue: the pivotal slide of this video. The written guide carries figure F5. "
           "Reveal the phases one at a time.")

# The module's emotional peak — the only full-colour punch block in the deck.
block(prs, 'Two rules for the order',
      ['Every prerequisite is in place before a service opens to the public.',
       'Citizens and businesses see a positive result in every phase.'],
      'PAERA 3.3.3: target the foundational elements — and show early results.',
      T,
      "VO: PAERA adds two rules. Every internal prerequisite must be in place before a service "
      "opens to the public. And citizens and businesses should always see a positive result. PAERA explains why both matter. Nobody wants to fund a digital identity while no "
      "service uses it, and no ministry can build personal services without one. The answer is "
      "to do two things at once. Put the foundational elements in place, and show early results "
      "that political leaders and the public can see. A service that opens before its "
      "foundations are ready fails in public, and the whole plan loses support.\n\n"
      "Production cue: the module's one full-colour block. Hold it a beat longer.",
      punch_fill=ITU_BLUE, punch_ink=WHITE)

rows_block(prs, 'What Estonia\'s shared blocks became',
           [('2000', 'The tax board\'s e-services.'),
            ('2001', 'X-Road, the data exchange layer.'),
            ('2002', 'The electronic identity and the digital signature.')],
           'The first services came before the shared blocks; once the blocks existed, they '
           'became part of the foundation.',
           T,
           "VO: Estonia's official record shows the order there. The tax board offered online tax "
           "returns in 2000. X-Road, the data exchange layer, followed in 2001, and the electronic "
           "identity with the digital signature in 2002. So in Estonia the first services came "
           "before the shared blocks. The lesson lies in what those blocks became once they "
           "existed. Estonia's own account names X-Road and the digital identity as part of the "
           "foundation that makes digital Estonia possible.\n\n"
           "On screen: the country name in plain typography only — no flags or emblems. Estonia "
           "is cited for what the shared blocks became, not as a case of identity coming before "
           "services.",
           numbered=False)

table_slide(prs, 'Progressa\'s order of building',
            ['Phase', 'What it brings', 'The result people see'], [1.3, 6.6, 4.0],
            [('1', 'Identity (PNIA) and a payment system (PayPro) already run: connect to them; '
                   'set up the decision body and the legal basis',
              'Online sign-in and fast payment, already offered to the public'),
             ('2', 'The ministry on Linkup; the registration service; the learner register, '
                   'brought forward; the Payments block on Linkup',
              'A parent registers a child once'),
             ('3', 'The learner register and registration in every district',
              'Every district registers learners the same way'),
             ('4', 'Every education service on the same blocks',
              'Parents and learners see their own records')],
            T,
            "VO: Here is the order for Progressa, the fictional country of these examples, on one "
            "page. Identity and payments already run there: PNIA, the identity authority, offers "
            "a sign-in, and PayPro, the payment provider, runs a fast-payment system. So phase one "
            "connects to them instead of building them, and sets up the decision body and the "
            "legal basis. Phase two brings the ministry of education onto Linkup, the data "
            "exchange. It builds the registration service, and brings the learner register "
            "forward because that service needs it. Its visible result: a parent registers a "
            "child once. Phase three takes the register to every district. Phase four lets "
            "parents and learners see their own records.",
            foot='Illustrative.', size=14)

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Identity and payments first. Data exchange and registration with the first priority "
      "services. The registers next. And one result people can see, early, so that support "
      "holds.",
      ('Order your planned services and blocks by what each depends on',
       'the order the dependencies allow, in four phases, with every service placed before a '
       'block it needs marked'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 3.3.3 and 5.7.1 to 5.7.5',
    'e-Estonia — \'e-Estonia story\' and \'Frequently asked questions — Story of e-Estonia\'',
])


# ================================================================ 1.10
T = '1.10 · The first proof for education: one service on four blocks'
MSG = ('The first proof for education is one service that registers a learner once: the '
       'sector\'s own registration service and learner register, using the country\'s identity '
       'and payment blocks through the data exchange layer.')
section('1.10', 'The first proof for education: one service on four blocks', MSG,
        "VO: A roadmap that promises everything proves nothing for years. Pick one service that "
        "shows the shared blocks serving real people, and prove it first. For education, that "
        "service is registering a learner once.")

rows_block(prs, 'Why one priority service',
           [('PAERA 5.7.3', 'High-priority services, prototyped quickly and rolled out in a few '
                            'months.'),
            ('PAERA Annex 1', 'A state registry needs two building blocks, Registration and '
                              'Digital Registry.')],
           'One registers the learner; the other keeps the record.',
           T,
           "VO: PAERA, the GovStack reference architecture, starts its second phase with high-priority services. Each is prototyped quickly, and then a local "
           "system integrator rolls it out in a few months. That gives the government its first "
           "practical experience of building fast on shared blocks. A learner registration is "
           "such a service. Families meet it when a child starts school, and every later "
           "education service needs its record. It needs two building blocks, Registration and a Digital Registry: one registers the learner; the other keeps the record.",
           numbered=False, bottom=5.3)

four_blocks(prs, 'One service, four blocks, on the data exchange layer', T,
            "VO: Here is the first proof. The registration service and the learner register are "
            "education's own, and the sector builds them. Identity and payments belong to the "
            "country, and the service uses them as they are. The data exchange layer joins them. "
            "The parent signs in through the national identity authority. The service registers "
            "the child, an officer approves, and the record is written to the learner register. "
            "The Payments block is set up beside the service and joined to the data exchange. A "
            "payment, such as a scholarship, can then reach a learner in the register.\n\n"
            "Production cue: the pivotal slide of this video and the spine of Modules 2 to 5. "
            "Reveal the service first, then education's two blocks, then the country's two, then "
            "the data exchange layer. Hold it a beat longer.")

table_slide(prs, 'Progressa\'s shortlist',
            ['Block', 'Owner', 'Specification and edition', 'Reused or built'],
            [1.9, 3.1, 3.2, 3.7],
            [('Registration', 'The ministry of education; PLR in charge of the registration',
              'GovStack Registration, default edition', 'Built for education'),
             ('Digital Registries', 'PLR, the learner registry',
              'GovStack Digital Registries, Version 3.0-alpha',
              'Built: the authoritative register behind PLR'),
             ('Identity', 'PNIA', 'GovStack Identity, Version 2.0',
              'Reused: the service signs people in through PNIA'),
             ('Payments', 'The government\'s Payments block, with PayPro behind its payer bank',
              'GovStack Payments, Version 3.0',
              'Set up and joined to Linkup; PayPro reused behind it'),
             ('Data exchange', 'PDGA, which operates Linkup',
              'GovStack Information Mediator, 1.1.1',
              'Reused: Linkup; the ministry joins as a member')],
            T,
            "VO: Here is the shortlist for Progressa, the fictional country of these examples, on "
            "one table. For each block it names the owner, the published specification with its "
            "edition, and whether the block is reused or built. Registration and the learner "
            "register are built for education. The register is set up behind PLR, the learner "
            "registry. PLR is a member of Linkup, the data exchange, but is not yet the "
            "authoritative register. Identity is reused: the service signs people in through "
            "PNIA, the identity authority. Linkup, operated by PDGA, the digital government "
            "authority, is reused, and the ministry joins it. The Payments block is set up and "
            "joined to Linkup, and it reaches PayPro, the payment provider, through its payer "
            "bank. Consent and messaging are cited where the service touches them, and neither "
            "is built.",
            foot='Consent (1.3.0) and Messaging are cited, not built. Illustrative.', size=12)

block(prs, 'Brought forward on purpose',
      ['PAERA 5.7.4 digitalises all main state registries in its third phase.',
       'This proof brings one register forward — the one its priority service needs. That choice '
       'is this course\'s own.'],
      'One priority service, proven on the country\'s foundation. Not the first phase of the '
      'national plan — and no replacement for it.',
      T,
      "VO: One choice must be stated openly. PAERA digitalises all main state registries in its third phase. This proof brings one register forward, the learner "
      "register, because its priority service cannot work without it. That choice is this course's own, not PAERA's. Present the proof to your minister as what it is: "
      "one priority service, proven on the country's foundation. It is not the first phase of "
      "the national plan, and it does not replace the plan.")

recap(MSG, T,
      PRACTICE_NOTE + "\n\n"
      "VO: Prove one service first: register a learner once, with the sector's own service and "
      "register, on the country's identity and payment blocks and its data exchange layer.",
      ('Draft your sector\'s shortlist of blocks for its first priority service',
       'a shortlist table with one row per block'))

sources_slide(prs, T, [
    'PAERA v1.0 — sections 5.7.3 and 5.7.4, and Annex 1 (A1.2.5)',
    'GovStack Building Block specifications — Registration (default edition), Digital Registries '
    '(Version 3.0-alpha; June 2026), Identity (Version 2.0; December 2025), Payments (Version '
    '3.0; December 2025)',
    'GovStack — Information Mediator (1.1.1), Consent (1.3.0), Messaging (messaging-23Q4.1)',
])


# ================================================================ Thank you
s = add_slide(prs, LAYOUT_THANKS)
notes(s, 'Closing slide for the combined deck. Individual videos end on their sources slide instead.')

# Self-check: the split spec's slide ranges depend on this count, and a helper that
# silently stops drawing shows up first as a slide with no voice-over.
assert len(prs.slides._sldIdLst) == 83, 'slide count changed — re-run the split with --infer-ranges'
assert all(sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip() for sl in prs.slides), \
    'every slide carries its voice-over in the notes'

# The practice box is the last thing on its slide; anything overlapping it clips on render.
for sl in prs.slides:
    for pb in [sh for sh in sl.shapes if sh.has_text_frame
               and sh.text_frame.text.startswith('Do this on your own sector')]:
        for sh in sl.shapes:
            if sh.shape_id == pb.shape_id or not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            assert sh.top + sh.height <= pb.top or sh.top >= pb.top + pb.height, \
                'shape overlaps the practice box: %r' % sh.text_frame.text[:60]

OUT = os.environ.get('OUT_PATH') or os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'videos', 'module_1', 'en', 'decks', 'KP3_M1_Deck_v0.1.pptx')
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print('slides:', len(prs.slides._sldIdLst))
print('saved', OUT)
