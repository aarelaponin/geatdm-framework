#!/usr/bin/env python3
# Shared content-shaped helpers for the KP3 Module 3–6 video decks (v0.1, 3 October 2026).
# Every generic helper, colour and layout index still comes from the kit's
# kp-deck-builder/scripts/deck_lib.py; this file only composes them, so that four module build
# scripts (build_kp3_module{3,4,5,6}_deck_v01.py) stay content only and share one layout.
# It carries the deck-side findings of KP3_M1-M2_Deck_v0.1_Review_2026-10-03.md that Modules 1
# and 2 have not yet taken:
#   - finding 10: tables at 15 pt or more (Modules 1–2 used 13 pt);
#   - finding 11: row type scales with the row count, and the block of rows sits centred in the
#     body area instead of being spread over a fixed band (no empty bottom third);
#   - finding 1 / 13: every storyboard stand-in carries the Module 2 footer
#     'Walkthrough: what a good run shows.' — one label, one place.
# The voice-over itself is never touched here: each build script passes the .js scriptBeats
# text verbatim (vo_diff.py proves it).
import os
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
KP_KIT = os.environ.get('KP_KIT') or os.path.join(_HERE, '..', '..', 'claude-marketplace', 'plugins',
                                                  'itu-giga-kp')
if not os.path.isdir(os.path.join(KP_KIT, 'skills', 'kp-deck-builder', 'scripts')):
    sys.exit("Set KP_KIT to the itu-giga-kp folder of your claude-marketplace clone (plugins/itu-giga-kp).")
sys.path.insert(0, os.path.join(KP_KIT, 'skills', 'kp-deck-builder', 'scripts'))
from pptx.util import Pt  # noqa: E402
from deck_lib import (  # noqa: E402
    TITLE_CARD_NOTE, hook_slide, practice_box,
    GREY, INK, ITU_BLUE_DARK, LIGHT, PANEL_GREY, WHITE,
    LAYOUT_BLUE, LAYOUT_THANKS, LAYOUT_WHITE,
    add_slide, big_slide, box, delete_template_slides, edit_agenda, edit_cover, footer, hline,
    notes, num_chip, open_template, panel, panel_text, section_slide, set_text, title)
from deck_diagrams import arrow, node  # noqa: E402

# The practice box is the video's only call to action and is never narrated (plan D5).
PRACTICE_NOTE = ('PRACTICE BOX (on-screen only — never read it, never paraphrase it, never point '
                 'at it). It replaces the narrated handoff: this video ends on the recap and the '
                 'Sources slide.')
STORYBOARD_FOOT = 'Walkthrough: what a good run shows.'
STORYBOARD_NOTE = ('On screen: the text stand-in for the demonstration segment, with the footer '
                   'marking the walkthrough, until the configurations it needs have passed their '
                   'checks and the segment is recorded.')

BODY_TOP, BODY_BOTTOM = 1.6, 6.55   # the content band under the headline, above the footer tag

# Row type and row height by row count (finding 11): few rows → larger type, tighter block.
_ROW_SCALE = {1: (26, 20, 1.5), 2: (24, 19, 1.4), 3: (22, 18, 1.3), 4: (20, 17, 1.12),
              5: (19, 16, 0.96), 6: (18, 15.5, 0.8)}


def _note(vo, cue=None, extra=None):
    parts = []
    if extra:
        parts.append(extra)
    parts.append('VO: ' + vo)
    if cue:
        parts.append('Production cue: ' + cue)
    return '\n\n'.join(parts)


def _line(s, text, y, size=16, bold=True, color=ITU_BLUE_DARK, italic=False, h=0.6):
    tb = box(s, 0.72, y, 11.9, h)
    set_text(tb.text_frame, [[(text, size, bold, color, italic)]])
    return tb


def _row_height(head, sub, hs, ss, rh):
    """Grow a row when its text will wrap past what the base height holds (11.9 in wide)."""
    def lines(t, pt):
        per = max(1, int(11.6 * 72 / (pt * 0.52)))   # ≈ characters per line of Arial at pt
        return -(-len(t) // per) if t else 0
    need = lines(head, hs) * hs * 1.25 / 72 + lines(sub, ss) * ss * 1.25 / 72 + 0.18
    return max(rh, round(need, 2))


class Deck:
    def __init__(self, module, out_name):
        self.m = module
        self.prs = open_template(os.environ.get('TEMPLATE'))
        self.out_name = out_name
        self.tag = ''
        self.msg = ''

    # ------------------------------------------------------------ cover and agenda
    def cover(self, **kw):
        edit_cover(self.prs, **kw)

    def agenda(self, **kw):
        edit_agenda(self.prs, **kw)
        delete_template_slides(self.prs, keep=2)

    # ------------------------------------------------------------ per video
    def video(self, code, name, message, hook, hook_vo):
        """Title card (silent cold open) + hook slide carrying the opener's VO."""
        self.tag = '%s · %s' % (code, name)
        self.msg = message
        section_slide(self.prs, 'MODULE %d · VIDEO %s' % (self.m, code), code, name, message,
                      'standalone video · voice-over on text slides', TITLE_CARD_NOTE)
        head, lines = hook
        hook_slide(self.prs, head, lines, self.tag, 'VO: ' + hook_vo)

    def _slide(self, head):
        s = add_slide(self.prs, LAYOUT_WHITE)
        title(s, head)
        return s

    def _end(self, s, vo, cue=None, extra=None):
        footer(s, self.tag)
        notes(s, _note(vo, cue, extra))
        return s

    def rows(self, head, rows, vo, closing=None, cue=None, lead=None, foot=None, numbered=False,
             extra=None, big=False):
        """Text rows: (headline, sub-line or ''). Type and height scale with the row count and the
        block is centred in the body band. `lead` = quoted lines above the rows (the
        specification's own words, italic); `closing` = the bold line the list lands on;
        `foot` = the small grey line under it (a source, an edition note, the storyboard)."""
        s = self._slide(head)
        n = len(rows)
        hs, ss, rh = _ROW_SCALE[min(n, 6)]
        if n > 6:
            rh = (BODY_BOTTOM - BODY_TOP - 0.3) / n
        if big:
            hs, ss, rh = 30, 20, 1.3
        heights = [_row_height(h, sb, hs, ss, rh) for h, sb in rows]
        lead_h = 0.5 * len(lead) + 0.15 if lead else 0
        close_h = 0.62 if closing else 0
        foot_h = 0.45 if foot else 0
        total = lead_h + sum(heights) + close_h + foot_h
        if total > BODY_BOTTOM - BODY_TOP:     # too tall: fall back to the band, step type down
            scale = (BODY_BOTTOM - BODY_TOP - lead_h - close_h - foot_h) / sum(heights)
            heights = [h * scale for h in heights]
            hs, ss = max(16, hs - 2), max(14, ss - 1.5)
            total = BODY_BOTTOM - BODY_TOP
        y = BODY_TOP + max(0, (BODY_BOTTOM - BODY_TOP - total) / 2 - 0.15)
        if lead:
            tb = box(s, 0.72, y, 11.9, lead_h)
            set_text(tb.text_frame, [[('“%s”' % q, 20, False, ITU_BLUE_DARK, True)] for q in lead],
                     space_after=Pt(4))
            y += lead_h
        for i, ((h, sb), rh_i) in enumerate(zip(rows, heights)):
            tx = 0.72
            if numbered:
                num_chip(s, 0.68, y + 0.07, i + 1)
                tx = 1.22
            tb = box(s, tx, y, 13.333 - tx - 0.6, rh_i - 0.05)
            paras = [[(h, hs, True, INK, False)]]
            if sb:
                paras.append([(sb, ss, False, GREY, False)])
            set_text(tb.text_frame, paras, space_after=Pt(3))
            y += rh_i
            if i < n - 1:
                hline(s, 0.68, y - 0.045, 11.9)
        if closing:
            _line(s, closing, y + 0.08, size=17 if len(closing) < 90 else 15.5)
            y += close_h
        if foot:
            _line(s, foot, y + 0.08, size=13, bold=False, color=GREY, italic=True, h=0.4)
        return self._end(s, vo, cue, extra)

    def storyboard(self, head, rows, vo, cue=None, numbered=True):
        """The text stand-in for a demonstration segment that has not been recorded."""
        return self.rows(head, rows, vo, cue=cue, foot=STORYBOARD_FOOT, numbered=numbered,
                         extra=STORYBOARD_NOTE)

    def table(self, head, cols, widths, rows, vo, cue=None, foot=None, size=None, extra=None):
        """A plain text table (finding 10: never under 15 pt). Header row in the accent colour,
        thin separators, no fills — text only, per the ITU guide. Height follows the content."""
        s = self._slide(head)
        size = size or (17 if len(rows) <= 3 else 16 if len(rows) <= 5 else 15)
        xs = [0.72]
        for w in widths[:-1]:
            xs.append(xs[-1] + w)
        hh = 0.5

        def rlines(row):
            m = 1
            for w, cell in zip(widths, row):
                per = max(1, int((w - 0.2) * 72 / (size * 0.5)))
                m = max(m, -(-len(cell) // per))
            return m
        rhs = [rlines(r) * size * 1.3 / 72 + 0.24 for r in rows]
        avail = BODY_BOTTOM - BODY_TOP - hh - (0.45 if foot else 0)
        if sum(rhs) > avail:
            k = avail / sum(rhs)
            rhs = [h * k for h in rhs]
        total = hh + sum(rhs) + (0.45 if foot else 0)
        top = BODY_TOP + max(0, (BODY_BOTTOM - BODY_TOP - total) / 2 - 0.15)
        for x, w, c in zip(xs, widths, cols):
            tb = box(s, x, top, w - 0.12, hh)
            set_text(tb.text_frame, [[(c, size, True, ITU_BLUE_DARK, False)]])
        hline(s, 0.72, top + hh, 11.9, color=ITU_BLUE_DARK, weight_pt=1.25)
        y = top + hh + 0.06
        for i, (row, rh) in enumerate(zip(rows, rhs)):
            for j, (x, w, cell) in enumerate(zip(xs, widths, row)):
                tb = box(s, x, y + 0.04, w - 0.12, rh - 0.08)
                set_text(tb.text_frame, [[(cell, size, j == 0, INK, False)]])
            y += rh
            if i < len(rows) - 1:
                hline(s, 0.72, y - 0.02, 11.9)
        if foot:
            _line(s, foot, y + 0.1, size=13, bold=False, color=GREY, italic=True, h=0.4)
        return self._end(s, vo, cue, extra)

    def panels(self, head, left, right, vo, closing=None, cue=None, foot=None, extra=None):
        """Two text boxes side by side: (heading, [lines]) each."""
        s = self._slide(head)
        nl = max(len(left[1]), len(right[1]))
        longest = max([len(t) for t in left[1] + right[1]] or [0])
        ch = min(4.2, 1.0 + nl * (0.5 if longest < 45 else 0.85))
        total = ch + (0.7 if closing else 0) + (0.45 if foot else 0)
        cy = BODY_TOP + max(0, (BODY_BOTTOM - BODY_TOP - total) / 2 - 0.15)
        cw = 5.9
        for i, ((ph, lines), fill) in enumerate(zip((left, right), (LIGHT, PANEL_GREY))):
            x = 0.72 + i * (cw + 0.23)
            panel(s, x, cy, cw, ch, fill)
            paras = [[(ph, 21, True, ITU_BLUE_DARK, False)]]
            paras += [[(ln, 18, False, INK, False)] for ln in lines]
            panel_text(s, x, cy, cw, ch, paras)
        y = cy + ch + 0.2
        if closing:
            _line(s, closing, y, size=18)
            y += 0.7
        if foot:
            _line(s, foot, y, size=13, bold=False, color=GREY, italic=True, h=0.4)
        return self._end(s, vo, cue, extra)

    def chain(self, head, steps, vo, closing=None, cue=None, foot=None, dark=None, extra=None,
              size=15):
        """Text boxes in a row joined by arrows. steps: (heading, [lines]); `dark` = index of the
        one box drawn in the accent colour."""
        s = self._slide(head)
        n = len(steps)
        gap = 0.42
        w = (11.9 - gap * (n - 1)) / n
        h = 2.5 if any(ln for _, ln in steps) else 1.6
        total = h + (0.75 if closing else 0) + (0.45 if foot else 0)
        y = BODY_TOP + max(0, (BODY_BOTTOM - BODY_TOP - total) / 2 - 0.15)
        for i in range(n - 1):
            x = 0.72 + (i + 1) * w + i * gap
            arrow(s, x + 0.04, y + h / 2, x + gap - 0.04, y + h / 2)
        for i, (name, lines) in enumerate(steps):
            d = dark == i
            node(s, 0.72 + i * (w + gap), y, w, h, name, lines,
                 fill=ITU_BLUE_DARK if d else LIGHT, ink=WHITE if d else INK,
                 head_ink=WHITE if d else ITU_BLUE_DARK, head_size=size + 1, size=size)
        y += h + 0.3
        if closing:
            _line(s, closing, y, size=18)
            y += 0.75
        if foot:
            _line(s, foot, y, size=13, bold=False, color=GREY, italic=True, h=0.4)
        return self._end(s, vo, cue, extra)

    def headings(self, head, cols, vo, closing=None, cue=None):
        """A one-page template shown as its column headings only (5.6's page for the next team)."""
        s = self._slide(head)
        n = len(cols)
        gap = 0.25
        w = (11.9 - gap * (n - 1)) / n
        y = 2.3
        for i, c in enumerate(cols):
            x = 0.72 + i * (w + gap)
            panel(s, x, y, w, 0.8, LIGHT)
            panel_text(s, x, y, w, 0.8, [[(c, 19, True, ITU_BLUE_DARK, False)]])
            for k in range(3):
                hline(s, x + 0.1, y + 1.35 + k * 0.6, w - 0.2)
        if closing:
            _line(s, closing, y + 3.3, size=18)
        return self._end(s, vo, cue)

    # ------------------------------------------------------------ recap and sources
    def recap(self, vo, practice):
        text = self.msg
        note = PRACTICE_NOTE + '\n\nVO: ' + vo
        if len(text) <= 150:
            return big_slide(self.prs, text, self.tag, note, practice=practice)
        s = add_slide(self.prs, LAYOUT_BLUE)
        tb = box(s, 1.1, 1.7, 6, 0.4)
        set_text(tb.text_frame, [[('IN ONE SENTENCE', 12, True, ITU_BLUE_DARK, False)]])
        tb = box(s, 1.1, 2.15, 11.1, 2.8)
        set_text(tb.text_frame, [[(text, 28 if len(text) <= 175 else 25 if len(text) <= 215 else 22,
                                   True, INK, False)]])
        practice_box(s, *practice, y=5.15)
        footer(s, self.tag, itu=True)
        notes(s, note)
        return s

    def sources(self, items):
        s = self._slide('Sources')
        n = sum(len(i) for i in items)
        size = 18 if n < 330 else 16 if n < 520 else 14
        tb = box(s, 0.72, 1.7, 11.8, 4.5)
        set_text(tb.text_frame, [[('•  ' + it, size, False, INK, False)] for it in items],
                 space_after=Pt(10))
        _line(s, 'Find the link in the description.', 6.35, size=15, h=0.5)
        footer(s, self.tag)
        notes(s, 'Sources slide, held ~5 seconds at the end of the video. Links are compiled into the '
                 'YouTube description per ITU convention; no URLs read aloud.')
        return s

    # ------------------------------------------------------------ finish
    def finish(self, expected):
        prs = self.prs
        s = add_slide(prs, LAYOUT_THANKS)
        notes(s, 'Closing slide for the combined deck. Individual videos end on their sources slide '
                 'instead.')
        n = len(prs.slides._sldIdLst)
        assert n == expected, 'slide count %d != %d — re-run the split with --infer-ranges' % (n, expected)
        assert all(sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip()
                   for sl in prs.slides), 'every slide carries its notes'
        for sl in prs.slides:    # nothing may overlap the practice box
            for pb in [sh for sh in sl.shapes if sh.has_text_frame
                       and sh.text_frame.text.startswith('Do this on your own sector')]:
                for sh in sl.shapes:
                    if sh.shape_id == pb.shape_id or not sh.has_text_frame \
                            or not sh.text_frame.text.strip():
                        continue
                    assert sh.top + sh.height <= pb.top or sh.top >= pb.top + pb.height, \
                        'shape overlaps the practice box: %r' % sh.text_frame.text[:60]
        out = os.environ.get('OUT_PATH') or os.path.join(
            _HERE, 'videos', 'module_%d' % self.m, 'en', 'decks', self.out_name)
        os.makedirs(os.path.dirname(out), exist_ok=True)
        prs.save(out)
        print('slides:', n)
        print('saved', out)
